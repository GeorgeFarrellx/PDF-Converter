# Version: 2.18
import os
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import zipfile
from datetime import datetime, timedelta, date

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# Drag & drop support
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except Exception as e:
    raise RuntimeError(
        "tkinterdnd2 is not installed. Install it with:\n"
        "  python -m pip install tkinterdnd2\n\n"
        f"Original error: {e}"
    )

from core import *  # noqa: F403
from core import _require_pdfplumber


def _read_app_version() -> str:
    version_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "VERSION.txt")
    try:
        with open(version_path, "r", encoding="utf-8") as f:
            return (f.read() or "").strip()
    except Exception:
        return ""


APP_VERSION = _read_app_version()



def _fmt_money(v) -> str:
    """Safely format a numeric value as GBP for logs/UI.

    Returns an empty string for None/blank. Never raises.
    """
    try:
        if v is None or v == "":
            return ""
        if isinstance(v, str):
            s = v.strip()
            if not s:
                return ""
            s = s.replace("£", "").replace(",", "")
            v = float(s)
        else:
            v = float(v)

        sign = "-" if v < 0 else ""
        v = abs(v)
        return f"{sign}£{v:,.2f}"
    except Exception:
        try:
            return str(v) if v is not None else ""
        except Exception:
            return ""


def show_reconciliation_popup(
    parent,
    output_path: str,
    recon_results: list[dict],
    coverage_period: str = "",
    continuity_results: list[dict] | None = None,
    audit_results: list[dict] | None = None,
    pre_save: bool = False,
    open_log_folder_callback=None,
    client_name: str = "",
):
    continuity_results = continuity_results or []
    audit_results = audit_results or []

    any_recon_warn = any((r.get("status") or "") != "OK" for r in (recon_results or []))
    any_cont_warn = any(
        not str((r.get("display_status") or r.get("status") or "")).strip().upper().startswith("OK")
        for r in (continuity_results or [])
    )
    any_audit_warn = any((a.get("status") or "") != "OK" for a in (audit_results or []))
    any_warn = any_recon_warn or any_cont_warn or any_audit_warn

    def _norm_date(v):
        try:
            if v is None or v == "":
                return None
            if hasattr(v, "to_pydatetime"):
                v = v.to_pydatetime()
            if hasattr(v, "date") and not isinstance(v, date):
                return v.date()
            return v
        except Exception:
            return None

    def _recon_sort_key(r):
        d = _norm_date(r.get("date_min"))
        d = d or _norm_date(r.get("date_max"))
        return (d or date.min, str(r.get("pdf") or ""))

    def _fmt_period(ps, pe) -> str:
        try:
            if ps and pe and hasattr(ps, "strftime") and hasattr(pe, "strftime"):
                return f"{ps.strftime('%d/%m/%Y')} - {pe.strftime('%d/%m/%Y')}"
        except Exception:
            return ""
        return ""

    def _safe_amount(value):
        try:
            if value is None:
                return None
            if isinstance(value, str):
                s = value.strip()
                if not s:
                    return None
                s = s.replace("£", "").replace(",", "")
                return float(s)
            return float(value)
        except Exception:
            return None

    def _fmt_money_or_na(value) -> str:
        text = _fmt_money(value)
        return text if text else "N/A"

    recon_results = sorted(list(recon_results or []), key=_recon_sort_key)
    period_by_pdf = {
        str(r.get("pdf") or ""): _fmt_period(r.get("period_start"), r.get("period_end"))
        for r in (recon_results or [])
    }

    win = tk.Toplevel(parent)
    base_bg = "#ffffff"
    header_bg = "#eef2f7"
    win.configure(bg=base_bg)

    win.title("Audit Checks (Warnings)" if any_warn else "Audit Checks")
    win.geometry("1200x560")
    win.minsize(1100, 560)
    win.resizable(True, True)
    win.lift()
    try:
        win.focus_set()
    except Exception:
        try:
            win.focus_force()
        except Exception:
            pass

    outer = tk.Frame(win, bg=base_bg)
    outer.pack(fill="both", expand=True, padx=10, pady=10)

    icon = "✖" if any_warn else "✔"
    icon_color = "#b00020" if any_warn else "#0b6e0b"

    head = tk.Frame(outer, bg=base_bg)
    head.pack(fill="x")

    tk.Label(head, text=icon, fg=icon_color, bg=base_bg, font=("Segoe UI", 18, "bold")).pack(side="left")

    title_text = "Audit Checks completed with warnings" if any_warn else "Audit Checks"
    tk.Label(head, text=title_text, bg=base_bg, font=("Segoe UI", 13, "bold")).pack(side="left", padx=(10, 0))

    path_row = tk.Frame(outer, bg=base_bg)
    path_row.pack(fill="x", pady=(10, 0))

    tk.Label(path_row, text="Output:", bg=base_bg).pack(side="left")
    tk.Label(path_row, text=output_path, fg="#333", bg=base_bg).pack(side="left", padx=(6, 0))

    PASS_SYMBOL = "✓"
    FAIL_SYMBOL = "✗"
    NA_SYMBOL = "—"

    if coverage_period:
        if any_warn:
            period_line = (
                f"The bank statements cover the period from {coverage_period} "
                "(however, some checks could not be completed or warnings were detected — see below)."
            )
        else:
            period_line = f"The bank statements cover the period from {coverage_period}."
    else:
        if any_warn:
            period_line = (
                "The bank statements cover the period: (unknown) "
                "(however, some checks could not be completed or warnings were detected — see below)."
            )
        else:
            period_line = "The bank statements cover the period: (unknown)."

    cn = str(client_name or "").strip()
    name_text = cn if cn else "(unknown)"

    info_bar = tk.Frame(outer, bg=base_bg)
    info_bar.pack(fill="x", pady=(8, 6))

    info_left = tk.Frame(info_bar, bg=base_bg)
    info_left.pack(side="left", fill="x", expand=True)
    tk.Label(info_left, text=name_text, bg=base_bg, font=("Segoe UI", 12, "bold"), anchor="w").pack(fill="x")
    tk.Label(
        info_left,
        text=period_line,
        bg=base_bg,
        fg="#333",
        font=("Segoe UI", 10),
        anchor="w",
        justify="left",
        wraplength=980,
    ).pack(fill="x", pady=(2, 0))

    audit_by_pdf = {
        str(a.get("pdf") or ""): a
        for a in (audit_results or [])
        if isinstance(a, dict)
    }

    style = ttk.Style(win)
    style.configure("SumHdr.TLabel", font=("Segoe UI", 10, "bold"))
    style.configure("SumFile.TLabel", font=("Segoe UI", 10))
    style.configure("SumPass.TLabel", foreground="#0b6e0b", font=("Segoe UI", 11, "bold"))
    style.configure("SumFail.TLabel", foreground="#b00020", font=("Segoe UI", 11, "bold"))
    style.configure("SumNA.TLabel", foreground="#666666", font=("Segoe UI", 11, "bold"))

    pdf_to_link_oks: dict[str, list[bool]] = {}
    for link in (continuity_results or []):
        if not isinstance(link, dict):
            continue
        st = str(link.get("display_status") or link.get("status") or "").strip().upper()
        link_ok = st.startswith("OK")
        prev_pdf = str(link.get("prev_pdf") or "")
        next_pdf = str(link.get("next_pdf") or "")
        if prev_pdf:
            pdf_to_link_oks.setdefault(prev_pdf, []).append(link_ok)
        if next_pdf:
            pdf_to_link_oks.setdefault(next_pdf, []).append(link_ok)

    file_display_width_chars = 42

    pdf_index = {str(r.get("pdf") or ""): idx for idx, r in enumerate(recon_results)}
    sorted_links = sorted(
        [link for link in continuity_results if isinstance(link, dict)],
        key=lambda l: (
            pdf_index.get(str(l.get("prev_pdf") or ""), 9999),
            pdf_index.get(str(l.get("next_pdf") or ""), 9999),
            str(l.get("prev_pdf") or ""),
            str(l.get("next_pdf") or ""),
        ),
    )

    selected_bg = "#cce8ff"
    row_bg_even = "#ffffff"
    row_bg_odd = "#f7f9fc"
    scroll_host = tk.Frame(outer, bg=base_bg)
    scroll_host.pack(fill="both", expand=True, pady=(6, 0))
    ysb = ttk.Scrollbar(scroll_host, orient="vertical")
    ysb.pack(side="right", fill="y")
    canvas = tk.Canvas(scroll_host, highlightthickness=0, background=base_bg)
    canvas.pack(side="left", fill="both", expand=True)
    canvas.configure(background=base_bg, yscrollcommand=ysb.set)
    ysb.configure(command=canvas.yview)

    content = tk.Frame(canvas, bg=base_bg)
    content_win_id = canvas.create_window((0, 0), window=content, anchor="nw")

    def _update_scrollregion(event=None):
        try:
            canvas.configure(scrollregion=canvas.bbox("all"))
        except Exception:
            pass

    content.bind("<Configure>", _update_scrollregion)

    def _sync_width(event=None):
        try:
            canvas.itemconfigure(content_win_id, width=canvas.winfo_width())
        except Exception:
            pass

    canvas.bind("<Configure>", _sync_width)

    def _event_in_this_popup(event) -> bool:
        try:
            w = event.widget
            return (w is not None) and (w.winfo_toplevel() == win)
        except Exception:
            return False

    def _scroll_units(units: int):
        try:
            canvas.yview_scroll(int(units), "units")
        except Exception:
            pass

    def _on_mousewheel_global(event):
        if not _event_in_this_popup(event):
            return
        try:
            delta = int(event.delta)
        except Exception:
            delta = 0
        if delta == 0:
            return "break"
        if sys.platform == "darwin":
            units = -1 if delta > 0 else 1
        else:
            step = int(delta / 80) if abs(delta) >= 80 else (1 if delta > 0 else -1)
            units = -step
        _scroll_units(units)
        return "break"

    def _on_mousewheel_linux_global(event):
        if not _event_in_this_popup(event):
            return
        try:
            if getattr(event, "num", None) == 4:
                _scroll_units(-1)
                return "break"
            if getattr(event, "num", None) == 5:
                _scroll_units(1)
                return "break"
        except Exception:
            pass
        return "break"

    win.bind_all("<MouseWheel>", _on_mousewheel_global, add="+")
    win.bind_all("<Button-4>", _on_mousewheel_linux_global, add="+")
    win.bind_all("<Button-5>", _on_mousewheel_linux_global, add="+")

    def _cleanup_binds():
        try:
            win.unbind_all("<MouseWheel>")
        except Exception:
            pass
        try:
            win.unbind_all("<Button-4>")
        except Exception:
            pass
        try:
            win.unbind_all("<Button-5>")
        except Exception:
            pass

    def _close():
        _cleanup_binds()
        win.destroy()

    card_border = "#d0d7de"

    def _make_card(parent, title: str):
        card = tk.Frame(parent, bg=base_bg, highlightbackground=card_border, highlightthickness=1)
        card.pack(fill="x", pady=(0, 10))
        hdr = tk.Frame(card, bg=header_bg)
        hdr.pack(fill="x")
        tk.Label(hdr, text=title, bg=header_bg, fg="#111", font=("Segoe UI", 10, "bold"), anchor="w", padx=8, pady=6).pack(fill="x")
        body = tk.Frame(card, bg=base_bg)
        body.pack(fill="x", padx=8, pady=6)
        return body
    win._selected_cell_text = ""
    win._selected_cell_widget = None
    table_data = {}

    def _select_cell(widget, value: str):
        try:
            prev = getattr(win, "_selected_cell_widget", None)
            if prev is not None and prev.winfo_exists():
                try:
                    prev_bg = getattr(prev, "_orig_bg", base_bg)
                    prev.configure(bg=prev_bg)
                except Exception:
                    pass
            try:
                widget.configure(bg=selected_bg)
            except Exception:
                pass
            win._selected_cell_widget = widget
            win._selected_cell_text = str(value or "")
            try:
                win.focus_set()
            except Exception:
                pass
            try:
                widget.focus_set()
            except Exception:
                pass
        except Exception:
            pass

    def _copy_tsv(text: str):
        try:
            win.clipboard_clear()
            win.clipboard_append(text)
            win.update_idletasks()
        except Exception:
            pass

    def _copy_selected_cell(event=None):
        try:
            val = str(getattr(win, "_selected_cell_text", "") or "")
            if not val:
                return "break"
            _copy_tsv(val)
        except Exception:
            pass
        return "break"

    def _cell_tsv(value: str) -> str:
        return str(value or "")
    def _copy_row_for(tbl_id, row_idx):
        try:
            if row_idx < 0 or tbl_id not in table_data:
                return
            headers = table_data[tbl_id]["headers"]
            row = table_data[tbl_id]["rows"][row_idx]
            tsv = "\t".join(map(_cell_tsv, headers)) + "\n" + "\t".join(map(_cell_tsv, row))
            _copy_tsv(tsv)
        except Exception:
            pass

    def _copy_table_for(tbl_id):
        try:
            if tbl_id not in table_data:
                return
            headers = table_data[tbl_id]["headers"]
            rows = table_data[tbl_id]["rows"]
            lines = ["\t".join(map(_cell_tsv, headers))]
            for r in rows:
                lines.append("\t".join(map(_cell_tsv, r)))
            _copy_tsv("\n".join(lines))
        except Exception:
            pass


    win.bind("<Control-c>", _copy_selected_cell, add="+")
    win.bind("<Control-C>", _copy_selected_cell, add="+")
    win.bind("<Command-c>", _copy_selected_cell, add="+")
    win.bind("<Command-C>", _copy_selected_cell, add="+")

    cell_menu = tk.Menu(win, tearoff=0)

    def _popup_cell_menu(event, widget, value):
        _select_cell(widget, value)
        cell_menu.delete(0, "end")
        tbl_id = getattr(widget, "_tbl_id", None)
        row_idx = getattr(widget, "_tbl_row", -1)
        cell_menu.add_command(label="Copy Cell", command=_copy_selected_cell)
        cell_menu.add_command(label="Copy Row", command=lambda t=tbl_id, r=row_idx: _copy_row_for(t, r))
        cell_menu.add_command(label="Copy Table", command=lambda t=tbl_id: _copy_table_for(t))
        try:
            cell_menu.tk_popup(event.x_root, event.y_root)
        finally:
            try:
                cell_menu.grab_release()
            except Exception:
                pass

    header_font = ("Segoe UI", 10, "bold")
    cell_font = ("Segoe UI", 10)
    symbol_font = ("Segoe UI", 11, "bold")
    pass_green = "#0b6e0b"
    fail_red = "#b00020"
    na_grey = "#666666"

    def _tbl_cell(
        parent,
        text,
        row,
        col,
        header=False,
        fg=None,
        anchor="w",
        width=None,
        font=None,
        justify="left",
        tbl_id=None,
        data_row_idx=-1,
        data_col_idx=-1,
        bind_wheel=None,
    ):
        chosen_bg = header_bg if header else (row_bg_odd if (row % 2 == 1) else row_bg_even)
        lbl = tk.Label(
            parent,
            text=text,
            bd=1,
            relief="solid",
            bg=chosen_bg,
            fg=(fg if fg is not None else "#000000"),
            anchor=anchor,
            width=width,
            justify=justify,
            font=(font or (header_font if header else cell_font)),
            padx=3,
            pady=1,
        )
        lbl._orig_bg = chosen_bg
        lbl._tbl_id = tbl_id
        lbl._tbl_row = data_row_idx
        lbl._tbl_col = data_col_idx
        lbl.grid(row=row, column=col, sticky="nsew")
        lbl.bind("<Button-1>", lambda e, w=lbl, t=text: _select_cell(w, t), add="+")
        lbl.bind("<Button-2>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        lbl.bind("<Button-3>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        lbl.bind("<Control-Button-1>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        return lbl

    def _tbl_merged_row(
        parent,
        text,
        row,
        colspan,
        tbl_id=None,
        data_row_idx=-1,
        data_col_idx=0,
        bg="#fff4ce",
    ):
        lbl = tk.Label(
            parent,
            text=text,
            bd=1,
            relief="solid",
            bg=bg,
            fg="#000000",
            anchor="center",
            justify="center",
            font=cell_font,
            padx=3,
            pady=1,
        )
        lbl._orig_bg = bg
        lbl._tbl_id = tbl_id
        lbl._tbl_row = data_row_idx
        lbl._tbl_col = data_col_idx
        lbl.grid(row=row, column=0, columnspan=colspan, sticky="nsew")
        lbl.bind("<Button-1>", lambda e, w=lbl, t=text: _select_cell(w, t), add="+")
        lbl.bind("<Button-2>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        lbl.bind("<Button-3>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        lbl.bind("<Control-Button-1>", lambda e, w=lbl, t=text: _popup_cell_menu(e, w, t), add="+")
        return lbl

    audit_body = _make_card(content, "Audit Summary")

    audit_tbl = ttk.Frame(audit_body)
    audit_tbl.pack(fill="x")

    headers = ["File", "Reconciliation", "Continuity", "Balance Walk", "Row Shape"]
    table_data["audit"] = {"headers": headers, "rows": []}
    for c, title in enumerate(headers):
        _tbl_cell(audit_tbl, title, 0, c, header=True, tbl_id="audit", data_row_idx=-1, data_col_idx=c)

    for row_idx, r in enumerate(recon_results, start=1):
        status = str(r.get("status") or "").strip()
        if status == "OK":
            recon_symbol = PASS_SYMBOL
            recon_fg = pass_green
        elif status == "Mismatch":
            recon_symbol = FAIL_SYMBOL
            recon_fg = fail_red
        elif status in ("Statement balances not found", "Not supported by parser") or "NOT CHECKED" in status.upper():
            recon_symbol = NA_SYMBOL
            recon_fg = na_grey
        else:
            recon_symbol = FAIL_SYMBOL
            recon_fg = fail_red

        pdf = str(r.get("pdf") or "")
        if len(pdf) > file_display_width_chars:
            pdf_disp = pdf[: file_display_width_chars - 1] + "…"
        else:
            pdf_disp = pdf

        link_oks = pdf_to_link_oks.get(pdf, [])
        if not link_oks:
            continuity_symbol = NA_SYMBOL
            continuity_fg = na_grey
        elif all(link_oks):
            continuity_symbol = PASS_SYMBOL
            continuity_fg = pass_green
        else:
            continuity_symbol = FAIL_SYMBOL
            continuity_fg = fail_red

        a = audit_by_pdf.get(pdf, {})

        bw_status = str(a.get("balance_walk_status") or "").strip()
        if not bw_status or bw_status.upper() == "NOT CHECKED":
            bw_symbol = NA_SYMBOL
            bw_fg = na_grey
        elif bw_status == "OK":
            bw_symbol = PASS_SYMBOL
            bw_fg = pass_green
        else:
            bw_symbol = FAIL_SYMBOL
            bw_fg = fail_red

        rs_status = str(a.get("row_shape_status") or "").strip()
        if not rs_status or rs_status.upper() == "NOT CHECKED":
            rs_symbol = NA_SYMBOL
            rs_fg = na_grey
        elif rs_status == "OK":
            rs_symbol = PASS_SYMBOL
            rs_fg = pass_green
        else:
            rs_symbol = FAIL_SYMBOL
            rs_fg = fail_red

        audit_row = [pdf_disp, recon_symbol, continuity_symbol, bw_symbol, rs_symbol]
        table_data["audit"]["rows"].append(audit_row)
        data_row_idx = row_idx - 1
        _tbl_cell(
            audit_tbl,
            pdf_disp,
            row_idx,
            0,
            width=file_display_width_chars,
            anchor="w",
            justify="left",
            tbl_id="audit",
            data_row_idx=data_row_idx,
            data_col_idx=0,
        )
        _tbl_cell(
            audit_tbl,
            recon_symbol,
            row_idx,
            1,
            fg=recon_fg,
            anchor="center",
            width=12,
            font=symbol_font,
            justify="center",
            tbl_id="audit",
            data_row_idx=data_row_idx,
            data_col_idx=1,
        )
        _tbl_cell(
            audit_tbl,
            continuity_symbol,
            row_idx,
            2,
            fg=continuity_fg,
            anchor="center",
            width=12,
            font=symbol_font,
            justify="center",
            tbl_id="audit",
            data_row_idx=data_row_idx,
            data_col_idx=2,
        )
        _tbl_cell(
            audit_tbl,
            bw_symbol,
            row_idx,
            3,
            fg=bw_fg,
            anchor="center",
            width=12,
            font=symbol_font,
            justify="center",
            tbl_id="audit",
            data_row_idx=data_row_idx,
            data_col_idx=3,
        )
        _tbl_cell(
            audit_tbl,
            rs_symbol,
            row_idx,
            4,
            fg=rs_fg,
            anchor="center",
            width=12,
            font=symbol_font,
            justify="center",
            tbl_id="audit",
            data_row_idx=data_row_idx,
            data_col_idx=4,
        )

    recon_body = _make_card(content, "Reconciliation")

    recon_tbl = ttk.Frame(recon_body)
    recon_tbl.pack(fill="x")

    recon_headers = [
        "File",
        "Credits",
        "Debits",
        "Total Txns",
        "Starting Balance",
        "Net Movement",
        "Calculated End",
        "Statement End",
        "Difference",
    ]
    table_data["recon"] = {"headers": recon_headers, "rows": []}
    for c, title in enumerate(recon_headers):
        header_anchor = "w" if c == 0 else "center"
        _tbl_cell(
            recon_tbl,
            title,
            0,
            c,
            header=True,
            anchor=header_anchor,
            justify="left" if c == 0 else "center",
            tbl_id="recon",
            data_row_idx=-1,
            data_col_idx=c,
        )

    recon_file_width_chars = 32
    for row_idx, r in enumerate(recon_results, start=1):
        pdf = str(r.get("pdf") or "")
        if len(pdf) > recon_file_width_chars:
            pdf_disp = pdf[: recon_file_width_chars - 1] + "…"
        else:
            pdf_disp = pdf

        transactions = r.get("transactions")
        if isinstance(transactions, list):
            txn_count = len(transactions)
        else:
            txn_count = 0
            for count_key in ("txn_count", "transaction_count"):
                count_val = r.get(count_key)
                if count_val is None:
                    continue
                try:
                    txn_count = int(count_val)
                    break
                except Exception:
                    continue

        credit_count = 0
        credit_total = 0.0
        debit_count = 0
        debit_total_abs = 0.0
        net_total = None

        if isinstance(transactions, list):
            for txn in transactions:
                if not isinstance(txn, dict):
                    continue
                amount = _safe_amount(txn.get("Amount"))
                if amount is None:
                    continue
                net_total = (net_total or 0.0) + amount
                if amount > 0:
                    credit_count += 1
                    credit_total += amount
                elif amount < 0:
                    debit_count += 1
                    debit_total_abs += abs(amount)
        else:
            net_from_result = _safe_amount(r.get("sum_amounts"))
            if net_from_result is not None:
                net_total = net_from_result

        start_balance = _safe_amount(r.get("start_balance"))
        end_balance = _safe_amount(r.get("end_balance"))
        difference = _safe_amount(r.get("difference"))
        expected_end = _safe_amount(r.get("expected_end"))

        if expected_end is not None:
            calculated_end = expected_end
        elif start_balance is not None and net_total is not None:
            calculated_end = start_balance + net_total
        else:
            calculated_end = None

        credits_cell = f"{credit_count}\n{_fmt_money_or_na(credit_total)}"
        debits_cell = f"{debit_count}\n{_fmt_money_or_na(debit_total_abs)}"
        row_values = [
            pdf_disp,
            credits_cell,
            debits_cell,
            txn_count,
            _fmt_money_or_na(start_balance),
            _fmt_money_or_na(net_total),
            _fmt_money_or_na(calculated_end),
            _fmt_money_or_na(end_balance),
            _fmt_money_or_na(difference),
        ]

        table_data["recon"]["rows"].append(row_values)
        data_row_idx = row_idx - 1
        for c, value in enumerate(row_values):
            anchor = "w" if c == 0 else "center"
            _tbl_cell(
                recon_tbl,
                value,
                row_idx,
                c,
                width=recon_file_width_chars if c == 0 else None,
                anchor=anchor,
                justify="left" if c == 0 else "center",
                tbl_id="recon",
                data_row_idx=data_row_idx,
                data_col_idx=c,
            )

    cont_body = _make_card(content, "Continuity")

    cont_tbl = ttk.Frame(cont_body)
    cont_tbl.pack(fill="x")

    cont_headers = [
        "File 1",
        "File 2",
        "Period 1",
        "Period 2",
        "File 1 End Balance",
        "File 2 Start Balance",
        "Status",
    ]
    table_data["cont"] = {"headers": cont_headers, "rows": []}
    for c, title in enumerate(cont_headers):
        _tbl_cell(cont_tbl, title, 0, c, header=True, tbl_id="cont", data_row_idx=-1, data_col_idx=c)

    file_link_width_chars = 28
    ui_row = 1
    for link in sorted_links:
        prev_pdf = str(link.get("prev_pdf") or "")
        next_pdf = str(link.get("next_pdf") or "")
        prev_end = link.get("prev_end")
        next_start = link.get("next_start")
        st = str(link.get("display_status") or link.get("status") or "").strip()
        st_upper = st.upper()

        if prev_end is None or next_start is None:
            status_text = "Not checked"
            status_style = "SumNA.TLabel"
        elif st_upper.startswith("OK"):
            status_text = "Match"
            status_style = "SumPass.TLabel"
        elif st_upper.startswith("MISMATCH"):
            status_text = "Mismatch"
            status_style = "SumFail.TLabel"
        else:
            status_text = st or "Not checked"
            if status_text == "Not checked":
                status_style = "SumNA.TLabel"
            else:
                status_style = "SumNA.TLabel"

        if len(prev_pdf) > file_link_width_chars:
            prev_pdf_disp = prev_pdf[: file_link_width_chars - 1] + "…"
        else:
            prev_pdf_disp = prev_pdf

        if len(next_pdf) > file_link_width_chars:
            next_pdf_disp = next_pdf[: file_link_width_chars - 1] + "…"
        else:
            next_pdf_disp = next_pdf

        row_values = [
            prev_pdf_disp,
            next_pdf_disp,
            period_by_pdf.get(prev_pdf, ""),
            period_by_pdf.get(next_pdf, ""),
            _fmt_money(prev_end) if prev_end is not None else "N/A",
            _fmt_money(next_start) if next_start is not None else "N/A",
        ]

        cont_row = row_values + [status_text]
        table_data["cont"]["rows"].append(cont_row)
        data_row_idx = len(table_data["cont"]["rows"]) - 1
        for c, value in enumerate(row_values):
            _tbl_cell(cont_tbl, value, ui_row, c, anchor="w", tbl_id="cont", data_row_idx=data_row_idx, data_col_idx=c)

        if status_style == "SumPass.TLabel":
            status_fg = pass_green
        elif status_style == "SumFail.TLabel":
            status_fg = fail_red
        else:
            status_fg = na_grey
        _tbl_cell(
            cont_tbl,
            status_text,
            ui_row,
            6,
            fg=status_fg,
            anchor="w",
            tbl_id="cont",
            data_row_idx=data_row_idx,
            data_col_idx=6,
        )
        ui_row += 1

        mf = link.get("missing_from")
        mt = link.get("missing_to")
        if mf is not None and mt is not None and hasattr(mf, "strftime") and hasattr(mt, "strftime"):
            gap_text = f"Suspected missing statement(s): {mf.strftime('%d/%m/%Y')} - {mt.strftime('%d/%m/%Y')}"
            table_data["cont"]["rows"].append([gap_text, "", "", "", "", "", ""])
            gap_data_row_idx = len(table_data["cont"]["rows"]) - 1
            _tbl_merged_row(
                cont_tbl,
                gap_text,
                ui_row,
                colspan=7,
                tbl_id="cont",
                data_row_idx=gap_data_row_idx,
                data_col_idx=0,
            )
            ui_row += 1

    bw_body = _make_card(content, "Balance Walk")

    bw_tbl = ttk.Frame(bw_body)
    bw_tbl.pack(fill="x")

    bw_headers = ["File", "Status", "Summary"]
    table_data["bw"] = {"headers": bw_headers, "rows": []}
    for c, title in enumerate(bw_headers):
        _tbl_cell(bw_tbl, title, 0, c, header=True, anchor="w", tbl_id="bw", data_row_idx=-1, data_col_idx=c)

    for row_idx, r in enumerate(recon_results, start=1):
        pdf = str(r.get("pdf") or "")
        if len(pdf) > recon_file_width_chars:
            pdf_disp = pdf[: recon_file_width_chars - 1] + "…"
        else:
            pdf_disp = pdf

        a = audit_by_pdf.get(pdf, {})
        bw_status = str(a.get("balance_walk_status") or "NOT CHECKED").strip()
        bw_summary = str(a.get("balance_walk_summary") or "").strip()

        if not bw_status or bw_status.upper() == "NOT CHECKED":
            bw_status_text = "NOT CHECKED"
            bw_status_fg = na_grey
        elif bw_status == "OK":
            bw_status_text = bw_status
            bw_status_fg = pass_green
        else:
            bw_status_text = bw_status
            bw_status_fg = fail_red

        bw_row = [pdf_disp, bw_status_text, bw_summary]
        table_data["bw"]["rows"].append(bw_row)
        data_row_idx = row_idx - 1
        _tbl_cell(bw_tbl, pdf_disp, row_idx, 0, width=recon_file_width_chars, anchor="w", tbl_id="bw", data_row_idx=data_row_idx, data_col_idx=0)
        _tbl_cell(bw_tbl, bw_status_text, row_idx, 1, fg=bw_status_fg, anchor="w", tbl_id="bw", data_row_idx=data_row_idx, data_col_idx=1)
        _tbl_cell(bw_tbl, bw_summary, row_idx, 2, anchor="w", tbl_id="bw", data_row_idx=data_row_idx, data_col_idx=2)

    rs_body = _make_card(content, "Row Shape Sanity")

    rs_tbl = ttk.Frame(rs_body)
    rs_tbl.pack(fill="x")

    rs_headers = ["File", "Status", "Summary"]
    table_data["rs"] = {"headers": rs_headers, "rows": []}
    for c, title in enumerate(rs_headers):
        _tbl_cell(rs_tbl, title, 0, c, header=True, anchor="w", tbl_id="rs", data_row_idx=-1, data_col_idx=c)

    for row_idx, r in enumerate(recon_results, start=1):
        pdf = str(r.get("pdf") or "")
        if len(pdf) > recon_file_width_chars:
            pdf_disp = pdf[: recon_file_width_chars - 1] + "…"
        else:
            pdf_disp = pdf

        a = audit_by_pdf.get(pdf, {})
        rs_status = str(a.get("row_shape_status") or "NOT CHECKED").strip()
        rs_summary = str(a.get("row_shape_summary") or "").strip()

        if not rs_status or rs_status.upper() == "NOT CHECKED":
            rs_status_text = "NOT CHECKED"
            rs_status_fg = na_grey
        elif rs_status == "OK":
            rs_status_text = rs_status
            rs_status_fg = pass_green
        else:
            rs_status_text = rs_status
            rs_status_fg = fail_red

        rs_row = [pdf_disp, rs_status_text, rs_summary]
        table_data["rs"]["rows"].append(rs_row)
        data_row_idx = row_idx - 1
        _tbl_cell(rs_tbl, pdf_disp, row_idx, 0, width=recon_file_width_chars, anchor="w", tbl_id="rs", data_row_idx=data_row_idx, data_col_idx=0)
        _tbl_cell(rs_tbl, rs_status_text, row_idx, 1, fg=rs_status_fg, anchor="w", tbl_id="rs", data_row_idx=data_row_idx, data_col_idx=1)
        _tbl_cell(rs_tbl, rs_summary, row_idx, 2, anchor="w", tbl_id="rs", data_row_idx=data_row_idx, data_col_idx=2)

    btn_row = ttk.Frame(win)
    btn_row.pack(fill="x", pady=(8, 10))
    create_support_bundle = getattr(parent, "create_support_bundle_zip", None)

    def _create_support_bundle():
        try:
            create_support_bundle()
        except Exception as e:
            messagebox.showerror("Support bundle error", str(e))

    if any_warn and callable(open_log_folder_callback):
        ttk.Button(btn_row, text="Open Log", command=open_log_folder_callback).pack(side="left", padx=(0, 8))

    close_btn = ttk.Button(btn_row, text="Close", command=_close)
    close_btn.pack(side="left")

    if callable(create_support_bundle):
        ttk.Button(btn_row, text="Create Support Bundle", command=_create_support_bundle).pack(side="left", padx=(8, 0))

    try:
        close_btn.focus_set()
        win.bind("<Return>", lambda e: _close())
        win.bind("<Escape>", lambda e: _close())
    except Exception:
        pass

    win.protocol("WM_DELETE_WINDOW", _close)

    parent.wait_window(win)
    return True


# ----------------------------
# GUI
# ----------------------------


class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()

        self.title("PDF Converter")
        self.geometry("760x520")
        self.minsize(720, 500)

        self.selected_files: list[str] = []
        self.bank_var = tk.StringVar(value="Select bank...")
        self._bank_set_by_autodetect = False
        self.output_folder_var = tk.StringVar(value=DEFAULT_OUTPUT_FOLDER)
        self.status_var = tk.StringVar(value="Ready.")
        self.auto_detect_var = tk.BooleanVar(value=True)

        self.last_report_data = None
        self.last_excel_data = None
        self.last_saved_output_path = None
        self._zip_temp_base_dir = ""
        self._zip_extracted_file_to_dir: dict[str, str] = {}

        self._build_ui()
        self._wire_dnd()
        self.protocol("WM_DELETE_WINDOW", self._on_app_close)

    def _build_ui(self):
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)

        bank_row = ttk.Frame(root)
        bank_row.pack(fill="x")

        bank_options = list(BANK_OPTIONS)
        if "Select bank..." in bank_options:
            bank_options = ["Select bank..."] + sorted(
                [opt for opt in bank_options if opt != "Select bank..."],
                key=str.lower,
            )
        else:
            bank_options = sorted(bank_options, key=str.lower)

        ttk.Label(bank_row, text="Bank:").pack(side="left")
        self.bank_combo = ttk.Combobox(
            bank_row,
            textvariable=self.bank_var,
            values=bank_options,
            state="readonly",
            width=20,
        )
        self.bank_combo.pack(side="left", padx=(8, 16))
        self.bank_combo.bind("<<ComboboxSelected>>", self._on_bank_selected, add="+")

        ttk.Checkbutton(
            bank_row,
            text="Auto-Detect Bank",
            variable=self.auto_detect_var,
        ).pack(side="left")

        ttk.Label(root, text="Drag & drop PDF statements here:").pack(anchor="w", pady=(14, 6))

        self.drop_box = tk.Listbox(root, height=10)
        self.drop_box.pack(fill="both", expand=False)
        self.drop_box.insert("end", "Drop PDFs here, or click 'Browse'.")

        btn_row = ttk.Frame(root)
        btn_row.pack(fill="x", pady=(10, 0))

        ttk.Button(btn_row, text="Browse", command=self.browse_pdfs).pack(side="left")
        ttk.Button(btn_row, text="Remove Selected", command=self.remove_selected).pack(side="left", padx=8)
        ttk.Button(btn_row, text="Clear List", command=self.clear_list).pack(side="left", padx=8)

        ttk.Separator(root, orient="horizontal").pack(fill="x", pady=(10, 10))

        run_row = ttk.Frame(root)
        run_row.pack(fill="x", pady=(12, 0))

        self.run_btn = ttk.Button(run_row, text="Convert", command=self.run_parser)
        self.run_btn.pack(side="left")

        self.cleanup_btn = ttk.Button(run_row, text="Clean", command=self.clean_up)
        self.cleanup_btn.pack(side="left", padx=10)

        ttk.Separator(root).pack(fill="x", pady=(18, 8))
        ttk.Label(root, text="Status:").pack(anchor="w")
        self.progress = ttk.Progressbar(root, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(4, 2))
        self.status_label = ttk.Label(root, textvariable=self.status_var)
        self.status_label.pack(anchor="w", pady=(2, 0))

        post_row = ttk.Frame(root)
        post_row.pack(fill="x", pady=(12, 0))

        self.show_checks_btn = ttk.Button(post_row, text="Show Checks", command=self.show_last_checks)
        self.show_checks_btn.pack(side="left")

        self.save_again_btn = ttk.Button(post_row, text="Save Output", command=self.save_last_output)
        self.save_again_btn.pack(side="left", padx=10)

        self.clear_logs_btn = ttk.Button(post_row, text="Clear Logs", command=self.clear_logs_folder)
        self.clear_logs_btn.pack(side="left", padx=10)

    def _wire_dnd(self):
        self.drop_box.drop_target_register(DND_FILES)
        self.drop_box.dnd_bind("<<Drop>>", self.on_drop)

    def set_status(self, msg: str):
        self.status_var.set(msg)
        try:
            self.update_idletasks()
        except Exception:
            pass

    def set_progress(self, completed: int, total: int):
        """Update progress bar based on completed PDFs / total PDFs."""
        try:
            total_i = int(total) if total not in (None, "") else 0
        except Exception:
            total_i = 0
        if total_i <= 0:
            total_i = 1
        try:
            completed_i = int(completed) if completed not in (None, "") else 0
        except Exception:
            completed_i = 0

        pct = (max(0, min(completed_i, total_i)) / total_i) * 100.0

        try:
            if hasattr(self, "progress") and self.progress is not None:
                self.progress["value"] = max(0.0, min(100.0, float(pct)))
        except Exception:
            pass

        try:
            self.update_idletasks()
        except Exception:
            pass

    def browse_pdfs(self):
        filepaths = filedialog.askopenfilenames(
            title="Select PDF statements or a ZIP",
            filetypes=[
                ("PDF and ZIP files", "*.pdf *.zip"),
                ("PDF files", "*.pdf"),
                ("ZIP files", "*.zip"),
            ],
        )
        if not filepaths:
            return

        selected_paths = list(filepaths)
        zip_paths = [p for p in selected_paths if p.lower().endswith(".zip")]
        pdf_paths = [p for p in selected_paths if p.lower().endswith(".pdf")]

        extracted_pdfs: list[str] = []
        for zip_path in zip_paths:
            try:
                extracted_pdfs.extend(self._extract_pdfs_from_zip(zip_path))
            except Exception as e:
                messagebox.showerror("ZIP Import", f"Failed to import ZIP: {e}")

        all_pdfs = pdf_paths + extracted_pdfs
        if not all_pdfs:
            messagebox.showwarning("No files", "No PDF files were selected or found in ZIP file(s).")
            return

        self.add_files(all_pdfs)

    def browse_zip(self):
        zip_path = filedialog.askopenfilename(
            title="Select ZIP file",
            filetypes=[("ZIP files", "*.zip")],
        )
        if not zip_path:
            return
        self.add_zip(zip_path)

    def browse_output_folder(self):
        folder = filedialog.askdirectory(title="Select output folder")
        if not folder:
            return
        self.output_folder_var.set(folder)

    def show_last_checks(self):
        if not self.last_report_data:
            messagebox.showwarning("Checks", "No checks to show yet. Convert first.")
            return

        data = self.last_report_data
        recon_results = data.get("recon_results") or []
        continuity_results = data.get("continuity_results") or []
        coverage_period = data.get("coverage_period") or ""
        any_warn = bool(data.get("any_warn"))
        output_path = (
            self.last_saved_output_path
            or data.get("output_xlsx_path")
            or "(Not saved yet)"
        )

        client_name = (data.get("client_name") or (self.last_excel_data or {}).get("client_name") or "")

        show_reconciliation_popup(
            self,
            output_path,
            recon_results,
            coverage_period=coverage_period,
            continuity_results=continuity_results,
            audit_results=data.get("audit_results") or [],
            pre_save=(output_path == "(Not saved yet)"),
            open_log_folder_callback=self.open_log_folder,
            client_name=client_name,
        )

    def _show_checks_text_popup(self, txt: str):
        win = tk.Toplevel(self)
        win.transient(self)
        win.grab_set()
        win.title("Checks Report")
        win.geometry("860x600")

        outer = ttk.Frame(win, padding=12)
        outer.pack(fill="both", expand=True)

        box = tk.Text(outer, wrap="word")
        box.pack(fill="both", expand=True)
        box.insert("1.0", txt or "")
        box.configure(state="disabled")

        btn_row = ttk.Frame(outer)
        btn_row.pack(fill="x", pady=(10, 0))
        ttk.Button(btn_row, text="Close", command=win.destroy).pack(side="right")

    def save_last_output(self):
        if not self.last_excel_data:
            messagebox.showwarning("Save", "Nothing to save yet. Convert first.")
            return

        data = self.last_excel_data
        transactions = data.get("transactions") or []
        client_name = data.get("client_name") or ""
        filename = data.get("filename") or "Transactions.xlsx"

        initial_dir = data.get("initial_dir") or ""
        if not initial_dir:
            folder = self.output_folder_var.get().strip()
            if folder:
                initial_dir = folder

        output_path = filedialog.asksaveasfilename(
            title="Save Excel file",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialdir=initial_dir or None,
            initialfile=filename,
        )

        if not output_path:
            self.set_status("Cancelled.")
            return

        try:
            ensure_folder(os.path.dirname(output_path))
        except Exception as e:
            messagebox.showerror("Save error", "Cannot create folder for output file:\n" + str(e))
            self.set_status("Failed")
            return

        try:
            self.output_folder_var.set(os.path.dirname(output_path))
        except Exception:
            pass

        hp_start = (self.last_excel_data or {}).get("statement_period_start")
        hp_end = (self.last_excel_data or {}).get("statement_period_end")

        self.set_status("Writing Excel...")
        save_transactions_to_excel(
            transactions,
            output_path,
            client_name=client_name,
            header_period_start=hp_start,
            header_period_end=hp_end,
        )

        self.last_saved_output_path = output_path
        self.set_progress(len(self.selected_files), max(1, len(self.selected_files)))

        # Determine whether the last run had warnings, so the status mirrors the main run result.
        any_warn = False
        try:
            any_warn = bool((self.last_report_data or {}).get("any_warn"))
        except Exception:
            any_warn = False

        if any_warn:
            self.set_status(f"Done with warnings. Output: {output_path}")
        else:
            self.set_status(f"Done. Output: {output_path}")

    def open_log_folder(self):
        try:
            ensure_folder(LOGS_DIR)
        except Exception as e:
            messagebox.showerror("Open log folder", f"Cannot access Logs folder:\n{e}")
            return

        log_path = os.path.abspath(LOGS_DIR)
        try:
            if sys.platform.startswith("win"):
                os.startfile(log_path)
            elif sys.platform == "darwin":
                subprocess.run(["open", log_path], check=False)
            else:
                subprocess.run(["xdg-open", log_path], check=False)
        except Exception as e:
            messagebox.showerror("Open log folder", f"Could not open log folder:\n{e}")

    def clear_logs_folder(self):
        confirmed = messagebox.askyesno(
            "Clear Logs",
            "This will permanently delete all files in the Logs folder, including support bundle ZIPs, crash logs, and recon logs. Continue?",
        )
        if not confirmed:
            return

        try:
            ensure_folder(LOGS_DIR)
        except Exception as e:
            messagebox.showerror("Clear Logs", f"Cannot access Logs folder:\n{e}")
            return

        deleted_files = 0
        deleted_dirs = 0
        failed_items = []

        for name in os.listdir(LOGS_DIR):
            if name == ".gitkeep":
                continue
            path = os.path.join(LOGS_DIR, name)
            try:
                if os.path.isdir(path):
                    shutil.rmtree(path, ignore_errors=False)
                    deleted_dirs += 1
                else:
                    os.remove(path)
                    deleted_files += 1
            except Exception:
                failed_items.append(name)

        summary = f"Deleted {deleted_files} file(s) and {deleted_dirs} folder(s)."
        if failed_items:
            preview = "\n".join(failed_items[:10])
            remaining = len(failed_items) - min(10, len(failed_items))
            if remaining > 0:
                preview += f"\n...and {remaining} more"
            summary += f"\n\nFailed to delete {len(failed_items)} item(s):\n{preview}"
            self.set_status("Logs cleared with some failures.")
        else:
            self.set_status("Logs cleared.")

        messagebox.showinfo("Clear Logs", summary)

    def create_support_bundle_zip(self):
        if not self.last_report_data:
            messagebox.showwarning("Support bundle", "No run data available. Run the parser first.")
            return

        data = self.last_report_data or {}

        learning_report_path = data.get("learning_report_path") or ""
        learning_report_inline = data.get("learning_report_inline") or ""
        learning_report_error = data.get("learning_report_error") or ""
        learning_report_generated = bool(data.get("learning_report_generated"))

        # Support bundle gating: treat any warnings/failures/NOT CHECKED/balances missing as issues.
        def _status_is_ok(s) -> bool:
            try:
                s = str(s or "").strip().upper()
            except Exception:
                return False
            return s.startswith("OK")

        def _text_has_issue_markers(s) -> bool:
            try:
                u = str(s or "").upper()
            except Exception:
                return False
            return (
                ("NOT CHECKED" in u)
                or ("BALANCES NOT FOUND" in u)
                or ("BALANCE NOT FOUND" in u)
                or ("MISMATCH" in u)
                or ("FAILED" in u)
            )

        def _missing_balances(rec: dict) -> bool:
            try:
                sb = rec.get("start_balance")
                eb = rec.get("end_balance")
                return (sb is None or sb == "") or (eb is None or eb == "")
            except Exception:
                return True

        has_issue = False
        issue_reasons = []

        # Fast path: if GUI already flagged warnings, allow bundle.
        try:
            has_issue = bool(data.get("any_warn"))
        except Exception:
            has_issue = False

        # Recon results: anything other than OK is an issue.
        if not has_issue:
            try:
                for r in (data.get("recon_results") or []):
                    st = (r.get("status") if isinstance(r, dict) else str(r or ""))
                    if (not _status_is_ok(st)) or _text_has_issue_markers(st):
                        has_issue = True
                        break
            except Exception:
                pass

        # Continuity results: NOT CHECKED is an issue (treat as error).
        if not has_issue:
            try:
                for c in (data.get("continuity_results") or []):
                    if isinstance(c, dict):
                        st = (c.get("display_status") or c.get("status") or "")
                        prev_pdf = c.get("prev_pdf") or ""
                        next_pdf = c.get("next_pdf") or ""
                    else:
                        st = str(c or "")
                        prev_pdf = ""
                        next_pdf = ""

                    if _text_has_issue_markers(st) or (not _status_is_ok(st)):
                        has_issue = True
                        if "NOT CHECKED" in str(st).upper():
                            issue_reasons.append(f"Continuity not checked: {prev_pdf} -> {next_pdf}")
                        break
            except Exception:
                pass

        # Audit results: anything other than OK is an issue.
        if not has_issue:
            try:
                for a in (data.get("audit_results") or []):
                    if not isinstance(a, dict):
                        continue
                    st = a.get("status") or ""
                    if (not _status_is_ok(st)) or _text_has_issue_markers(st):
                        has_issue = True
                        break
            except Exception:
                pass

        try:
            for a in (data.get("audit_results") or []):
                if not isinstance(a, dict):
                    continue
                st = a.get("status") or ""
                if _status_is_ok(st):
                    continue
                pdf = a.get("pdf") or "PDF"
                issue_reasons.append(
                    f"Audit: {pdf} | Balance Walk {a.get('balance_walk_status') or 'NOT CHECKED'}"
                    f" ({a.get('balance_walk_summary') or ''}) | Row Shape Sanity {a.get('row_shape_status') or 'WARN'}"
                    f" ({a.get('row_shape_summary') or ''})"
                )
        except Exception:
            pass

        # Explicit balance-missing gate: if any PDF lacks start/end balances, continuity may show "Not checked".
        # Treat this as an issue so support bundles include logs.
        try:
            for r in (data.get("recon_results") or []):
                if not isinstance(r, dict):
                    continue
                if _missing_balances(r):
                    has_issue = True
                    issue_reasons.append(f"{r.get('pdf') or 'PDF'}: balances not found")
        except Exception:
            pass

        if not has_issue:
            # Still allow bundle for verification/debugging, but wording is safe (no broken string literals).
            messagebox.showinfo(
                "Support bundle",
                "No reconciliation or continuity issues were detected for the last run based on stored results.\n"
                "You can still create a support bundle for verification/debugging.",
            )

        if not self.last_excel_data:
            messagebox.showerror(
                "Support bundle",
                "Cannot create support bundle because the Excel data for the last run is missing.",
            )
            return

        try:
            ensure_folder(LOGS_DIR)
        except Exception as e:
            messagebox.showerror("Support bundle", "Cannot access Logs folder:\n" + str(e))
            return

        excel_source = (data.get("output_xlsx_path") or self.last_saved_output_path or "")
        excel_source = excel_source if (excel_source and os.path.exists(excel_source)) else ""

        bundle_base = data.get("bundle_base")
        if not bundle_base:
            try:
                bundle_base = os.path.splitext(self.last_excel_data.get("filename") or "Transactions.xlsx")[0]
            except Exception:
                bundle_base = "RUN"

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_base = sanitize_filename(bundle_base) or "RUN"
        zip_base = re.sub(r"\s\d{2}\.\d{2}\.\d{2}\s*-\s*\d{2}\.\d{2}\.\d{2}$", "", safe_base).strip()
        if not zip_base:
            zip_base = safe_base
        zip_name = f"{zip_base}.zip"
        zip_path = make_unique_path(os.path.join(LOGS_DIR, zip_name))

        transactions = self.last_excel_data.get("transactions") or []
        client_name = self.last_excel_data.get("client_name") or ""
        hp_start = (self.last_excel_data or {}).get("statement_period_start")
        hp_end = (self.last_excel_data or {}).get("statement_period_end")

        temp_excel_path = ""
        created_temp_excel = False
        zip_created = False
        excel_creation_error = ""

        def _create_empty_support_excel(output_path: str, client_name_for_header: str = ""):
            try:
                from openpyxl import Workbook
            except Exception as e:
                _show_dependency_error(
                    "openpyxl is required for support bundle Excel output.\n\n"
                    "Install it with:\n"
                    "  python -m pip install openpyxl\n\n"
                    f"Original error: {e}"
                )
                return None

            wb = Workbook()
            ws = wb.active
            ws.title = "Transaction Data"

            headers = ["T/N", "Date", "Transaction Type", "Description", "Amount", "Balance", "Category"]
            ws.append(headers)

            left_text = (client_name_for_header or "").strip()
            center_text = "Transaction Data"
            right_text = ""
            for hdr in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
                hdr.left.text = left_text
                hdr.center.text = center_text
                hdr.right.text = right_text

            ensure_folder(os.path.dirname(output_path))
            wb.save(output_path)
            return output_path

        try:
            if not excel_source:
                temp_excel_name = f"SUPPORT EXCEL - {bundle_base} - {ts}.xlsx"
                temp_excel_path = make_unique_path(os.path.join(LOGS_DIR, temp_excel_name))
                try:
                    if transactions:
                        save_transactions_to_excel(
                            transactions,
                            temp_excel_path,
                            client_name=client_name,
                            header_period_start=hp_start,
                            header_period_end=hp_end,
                        )
                    else:
                        _create_empty_support_excel(temp_excel_path, client_name_for_header=client_name)
                    excel_source = temp_excel_path
                    created_temp_excel = True
                except Exception as e:
                    excel_source = ""
                    excel_creation_error = "".join(traceback.format_exception(type(e), e, e.__traceback__))

            recon_log_path = data.get("log_path") or ""
            log_path = recon_log_path
            support_log_path = ""
            log_exists = bool(log_path and os.path.exists(log_path))

            # If no log exists (common when continuity is NOT CHECKED), create a lightweight support log now.
            if not log_exists:
                support_log_name = f"{bundle_base} - support log - {ts}.txt"
                support_log_path = make_unique_path(os.path.join(LOGS_DIR, support_log_name))

                recon = data.get("recon_results") or []
                cont = data.get("continuity_results") or []
                audit = data.get("audit_results") or []

                lines = []
                lines.append(f"Support log generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
                lines.append(f"Bank: {(data.get('bank') or '').strip()}")
                lines.append(f"GUI: {APP_VERSION}")
                lines.append(f"Core: {APP_VERSION}")
                lines.append("")

                if issue_reasons:
                    lines.append("Detected issues:")
                    for r in issue_reasons:
                        lines.append(f"- {r}")
                    lines.append("")

                lines.append("Reconciliation summary:")
                for r in recon:
                    if not isinstance(r, dict):
                        lines.append(f"- {r}")
                        continue
                    pdf = r.get("pdf") or ""
                    st = r.get("status") or ""
                    sb = _fmt_money(r.get("start_balance"))
                    eb = _fmt_money(r.get("end_balance"))
                    ps = r.get("period_start")
                    pe = r.get("period_end")
                    per = ""
                    if ps and pe and hasattr(ps, "strftime") and hasattr(pe, "strftime"):
                        per = f"{ps.strftime('%d/%m/%Y')} - {pe.strftime('%d/%m/%Y')}"
                    lines.append(
                        f"- {pdf}: {st} | start={sb or '<missing>'} | end={eb or '<missing>'} | period={per or 'None'}"
                    )
                lines.append("")

                lines.append("Continuity summary:")
                for c in cont:
                    if isinstance(c, dict):
                        prev_pdf = c.get("prev_pdf") or ""
                        next_pdf = c.get("next_pdf") or ""
                        st = c.get("display_status") or c.get("status") or ""
                        lines.append(f"- {prev_pdf} -> {next_pdf}: {st}")
                    else:
                        lines.append(f"- {c}")
                lines.append("")

                lines.append("Audit summary:")
                for a in audit:
                    if not isinstance(a, dict):
                        lines.append(f"- {a}")
                        continue
                    pdf = a.get("pdf") or ""
                    lines.append(
                        f"- {pdf}: Balance Walk {a.get('balance_walk_status') or 'NOT CHECKED'}"
                        f" ({a.get('balance_walk_summary') or ''}) | Row Shape Sanity {a.get('row_shape_status') or 'WARN'}"
                        f" ({a.get('row_shape_summary') or ''}) | Overall {a.get('status') or 'WARN'}"
                    )
                lines.append("")

                with open(support_log_path, "w", encoding="utf-8") as f:
                    f.write("\n".join(lines).rstrip() + "\n")

                log_path = support_log_path
                log_exists = True
                try:
                    if self.last_report_data is not None:
                        self.last_report_data["log_path"] = log_path
                except Exception:
                    pass

            if (
                not learning_report_generated
                and not learning_report_path
                and not learning_report_inline
                and not learning_report_error
            ):
                try:
                    report_path, report_text, report_err = self.generate_learning_report(
                        reason="Support bundle", write_to_disk=False
                    )
                    learning_report_path = report_path or ""
                    learning_report_inline = report_text or ""
                    learning_report_error = report_err or ""
                except Exception as e:
                    learning_report_error = "".join(
                        traceback.format_exception(type(e), e, e.__traceback__)
                    )

            pdf_paths = list(data.get("source_pdfs") or [])

            used_names = set()

            def _unique_zip_name(filename: str) -> str:
                base, ext = os.path.splitext(filename)
                if not ext:
                    ext = ".pdf"
                candidate = f"{base}{ext}"
                if candidate.lower() not in used_names:
                    used_names.add(candidate.lower())
                    return candidate
                n = 2
                while True:
                    candidate = f"{base} ({n}){ext}"
                    if candidate.lower() not in used_names:
                        used_names.add(candidate.lower())
                        return candidate
                    n += 1

            snapshot_files = []
            missing_snapshot = []

            gui_path = os.path.abspath(__file__)
            if os.path.exists(gui_path):
                snapshot_files.append((gui_path, "CODE_SNAPSHOT/gui.py"))
            else:
                missing_snapshot.append("gui.py")

            core_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "core.py")
            if os.path.exists(core_path):
                snapshot_files.append((core_path, "CODE_SNAPSHOT/core.py"))
            else:
                missing_snapshot.append("core.py")

            parser_path = data.get("parser_file") or ""
            parser_basename = os.path.basename(parser_path) if parser_path else ""
            if parser_path and os.path.exists(parser_path):
                snapshot_files.append((parser_path, f"CODE_SNAPSHOT/{parser_basename}"))
            else:
                missing_snapshot.append(parser_basename or "parser file")

            with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                excel_arc = f"{safe_base}.xlsx"
                if excel_source and os.path.exists(excel_source):
                    zf.write(excel_source, arcname=excel_arc)
                elif excel_creation_error:
                    zf.writestr("EXCEL_CREATION_FAILED.txt", excel_creation_error)
                else:
                    zf.writestr(
                        "EXCEL_CREATION_FAILED.txt",
                        "Support bundle could not include an Excel file because no source Excel was available.\n",
                    )

                if log_exists:
                    zf.write(log_path, arcname="Reconciliation Log.txt")

                if learning_report_inline:
                    zf.writestr("Learning Report.txt", learning_report_inline)
                elif learning_report_path and os.path.exists(learning_report_path):
                    zf.write(learning_report_path, arcname="Learning Report.txt")
                elif learning_report_error:
                    zf.writestr("LEARNING_FAILED_INLINE.txt", learning_report_error)
                else:
                    lines = [
                        "Learning report was not created or could not be found on disk at bundle time.",
                        "If LEARNING_FAILED exists in Logs, it should be included.",
                    ]
                    zf.writestr("LEARNING_REPORT_MISSING.txt", "\n".join(lines).rstrip() + "\n")

                for p in pdf_paths:
                    if not p or not os.path.exists(p):
                        continue
                    base = os.path.basename(p) or "statement.pdf"
                    safe = sanitize_filename(base) or "statement.pdf"
                    unique = _unique_zip_name(safe)
                    arcname = "Source PDFs/" + unique
                    zf.write(p, arcname=arcname)

                for local_path, arcname in snapshot_files:
                    zf.write(local_path, arcname=arcname)

                if missing_snapshot:
                    zf.writestr(
                        "CODE_SNAPSHOT/_MISSING_FILES.txt",
                        "\n".join(sorted(set(missing_snapshot))) + "\n",
                    )

            zip_created = True

            for p in [recon_log_path, support_log_path, learning_report_path]:
                if p and os.path.exists(p):
                    try:
                        os.remove(p)
                    except Exception:
                        pass

            try:
                cutoff = datetime.now() - timedelta(minutes=2)
                for name in os.listdir(LOGS_DIR):
                    if not name.startswith("LEARNING - ") or not name.lower().endswith(".txt"):
                        continue
                    full_path = os.path.join(LOGS_DIR, name)
                    try:
                        mtime = datetime.fromtimestamp(os.path.getmtime(full_path))
                    except Exception:
                        continue
                    if mtime < cutoff:
                        continue
                    try:
                        os.remove(full_path)
                    except Exception:
                        pass
            except Exception:
                pass

            if created_temp_excel and temp_excel_path and os.path.exists(temp_excel_path):
                try:
                    os.remove(temp_excel_path)
                except Exception:
                    pass

            messagebox.showinfo("Support bundle created", "Support bundle created:\n" + zip_path)

        except Exception as e:
            if created_temp_excel and temp_excel_path and zip_created and os.path.exists(temp_excel_path):
                try:
                    os.remove(temp_excel_path)
                except Exception:
                    pass

            err = "".join(traceback.format_exception(type(e), e, e.__traceback__))
            messagebox.showerror("Support bundle error", f"{e}\n\nDetails:\n{err}")


    def generate_learning_report(
        self,
        reason: str | None = None,
        exception: Exception | None = None,
        write_to_disk: bool = False,
    ):
        pdfplumber = _require_pdfplumber(show_error=False)
        pdfplumber_available = pdfplumber is not None

        data = self.last_report_data or {}
        bank = (data.get("bank") or self.bank_var.get() or "").strip() or "Unknown"
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_name = f"LEARNING - {ts}.txt"

        source_pdfs = list(data.get("source_pdfs") or [])

        main_id = APP_VERSION

        autodetect_result = data.get("autodetect_first_pdf")
        parser_file = data.get("parser_file") or ""

        recon_results = data.get("recon_results") or []
        continuity_results = data.get("continuity_results") or []

        total_tx = 0
        date_min = None
        date_max = None
        try:
            if self.last_excel_data:
                txs = self.last_excel_data.get("transactions") or []
                total_tx = len(txs)
                dates = [t.get("Date") for t in txs if isinstance(t, dict) and t.get("Date")]
                if dates:
                    date_min = min(dates)
                    date_max = max(dates)
        except Exception:
            total_tx = 0
            date_min = None
            date_max = None

        def _norm_text_block(s: str) -> str:
            s = (s or "").replace("\r\n", "\n").replace("\r", "\n")
            lines = [ln.strip() for ln in s.split("\n")]
            out = []
            blank_run = 0
            for ln in lines:
                if not ln:
                    blank_run += 1
                    if blank_run <= 2:
                        out.append("")
                    continue
                blank_run = 0
                out.append(ln)
            return "\n".join(out).strip()

        def _page_snapshot(text: str) -> str:
            text = _norm_text_block(text)
            if not text:
                return ""
            lines = text.splitlines()
            if len(lines) > 80:
                snap = "\n".join(lines[:80])
            else:
                snap = text
            if len(snap) > 3000:
                snap = snap[:3000] + "\n...<truncated>"
            return snap

        def _fmt_date(v):
            try:
                if v is None or v == "":
                    return ""
                if hasattr(v, "to_pydatetime"):
                    v = v.to_pydatetime()
                if hasattr(v, "date") and isinstance(v, datetime):
                    v = v.date()
                if hasattr(v, "strftime"):
                    return v.strftime("%d/%m/%Y")
                return str(v)
            except Exception:
                return str(v)

        lines = []
        try:
            lines.append("LEARNING REPORT")
            lines.append("=" * 60)
            lines.append(f"Timestamp: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            if main_id:
                lines.append(f"Main: {main_id}")
            lines.append(f"Bank (selected): {bank}")
            if autodetect_result:
                lines.append(f"Auto-detect (first PDF): {autodetect_result}")
            if parser_file:
                lines.append(f"Parser file: {parser_file}")
            if reason:
                lines.append(f"Reason: {reason}")
            lines.append("")

            lines.append("PDF SUMMARY")
            lines.append("-" * 60)
            if not pdfplumber_available:
                lines.append("pdfplumber is not available; skipping PDF text extraction for this report.")
                lines.append("")

            for pdf_path in source_pdfs:
                lines.append(f"PDF: {os.path.basename(pdf_path)}")
                lines.append(f"Path: {pdf_path}")

                page_count = 0
                empty_pages = 0
                per_page_text = []

                if pdfplumber_available:
                    try:
                        with pdfplumber.open(pdf_path) as pdf:
                            page_count = len(pdf.pages)
                            for pi, page in enumerate(pdf.pages, start=1):
                                txt = ""
                                try:
                                    txt = page.extract_text() or ""
                                except Exception:
                                    txt = ""

                                snap_base = _norm_text_block(txt)
                                if len(snap_base) < 50:
                                    empty_pages += 1
                                per_page_text.append((pi, txt))

                    except Exception as e:
                        lines.append(f"Page count: (error: {e})")
                        lines.append("")
                        continue

                if pdfplumber_available:
                    lines.append(f"Page count: {page_count}")
                    mostly_empty = (page_count > 0 and (empty_pages / page_count) >= 0.7)
                    lines.append(
                        "Extracted text mostly empty: "
                        f"{'YES' if mostly_empty else 'NO'} ({empty_pages}/{page_count} pages low-text)"
                    )
                    lines.append("")

                    lines.append("Per-page text snapshot:")
                    for (pi, txt) in per_page_text:
                        snap = _page_snapshot(txt)
                        lines.append("")
                        lines.append(f"--- Page {pi} ---")
                        if snap:
                            lines.append(snap)
                        else:
                            lines.append("<no extracted text>")
                else:
                    lines.append("Page count: (skipped - pdfplumber unavailable)")
                    lines.append("Per-page text snapshot: (skipped - pdfplumber unavailable)")

                lines.append("")
                lines.append("-" * 60)
                lines.append("")

            lines.append("RUN SUMMARY")
            lines.append("-" * 60)
            lines.append(f"Total transactions: {total_tx}")
            if date_min and date_max:
                lines.append(f"Date range: {_fmt_date(date_min)} - {_fmt_date(date_max)}")
            lines.append("")

            lines.append("Statement balances found per PDF:")
            for r in recon_results:
                pdf = r.get("pdf") or ""
                sb = r.get("start_balance")
                eb = r.get("end_balance")
                sb_ok = "YES" if sb is not None and sb != "" else "NO"
                eb_ok = "YES" if eb is not None and eb != "" else "NO"
                lines.append(f"- {pdf}: start_found={sb_ok}, end_found={eb_ok}")
            lines.append("")

            lines.append("Reconciliation results:")
            for r in recon_results:
                pdf = r.get("pdf") or ""
                st = r.get("status") or ""
                diff = r.get("difference")
                line = f"- {pdf}: {st}"
                if diff is not None and diff != "":
                    try:
                        line += f" (diff {float(diff):.2f})"
                    except Exception:
                        line += f" (diff {diff})"
                lines.append(line)
            lines.append("")

            if continuity_results:
                lines.append("Continuity results:")
                for c in continuity_results:
                    prev_pdf = c.get("prev_pdf") or ""
                    next_pdf = c.get("next_pdf") or ""
                    st = c.get("status") or ""
                    diff = c.get("diff")
                    missing = ""
                    try:
                        mf = c.get("missing_from")
                        mt = c.get("missing_to")
                        if mf and mt and hasattr(mf, "strftime") and hasattr(mt, "strftime"):
                            missing = f" | missing {_fmt_date(mf)} - {_fmt_date(mt)}"
                    except Exception:
                        missing = ""

                    line = f"- {prev_pdf} -> {next_pdf}: {st}"
                    if diff is not None and diff != "":
                        try:
                            line += f" (diff {float(diff):.2f})"
                        except Exception:
                            line += f" (diff {diff})"
                    if missing:
                        line += missing
                    lines.append(line)
                lines.append("")

            if exception is not None:
                lines.append("EXCEPTION")
                lines.append("-" * 60)
                lines.append(f"Type: {type(exception).__name__}")
                lines.append(f"Message: {exception}")
                lines.append("")
                tb = "".join(traceback.format_exception(type(exception), exception, exception.__traceback__))
                lines.append(tb.rstrip())
                lines.append("")
        except Exception as e:
            err = "".join(traceback.format_exception(type(e), e, e.__traceback__))
            try:
                if self.last_report_data is None:
                    self.last_report_data = {}
                self.last_report_data["learning_report_error"] = err
            except Exception:
                pass
            return None, None, err

        report_text = "\n".join(lines).rstrip() + "\n"

        try:
            if self.last_report_data is None:
                self.last_report_data = {}
            self.last_report_data["learning_report_generated"] = True
        except Exception:
            pass

        if not write_to_disk:
            try:
                if self.last_report_data is None:
                    self.last_report_data = {}
                self.last_report_data["learning_report_inline"] = report_text
            except Exception:
                pass
            return None, report_text, None

        try:
            ensure_folder(LOGS_DIR)
            report_path = make_unique_path(os.path.join(LOGS_DIR, report_name))
            with open(report_path, "w", encoding="utf-8") as f:
                f.write(report_text)
        except Exception as e:
            err = "".join(traceback.format_exception(type(e), e, e.__traceback__))
            try:
                if self.last_report_data is None:
                    self.last_report_data = {}
                self.last_report_data["learning_report_inline"] = report_text
                self.last_report_data["learning_report_error"] = err
            except Exception:
                pass
            return None, report_text, err

        try:
            if self.last_report_data is None:
                self.last_report_data = {}
            self.last_report_data["learning_report_path"] = report_path
        except Exception:
            pass

        return report_path, report_text, None

    def clear_list(self):
        removed_paths = list(self.selected_files or [])
        self.selected_files = []
        self._cleanup_removed_zip_files(removed_paths)
        self._cleanup_all_zip_temp()
        self.drop_box.delete(0, "end")
        self.drop_box.insert(0, "Drop PDFs here, or click 'Browse'.")
        self._reset_bank_if_autodetected()
        self.set_status("Cleared file list.")

    def remove_selected(self):
        selected = list(self.drop_box.curselection())
        if not selected:
            return

        displayed = list(self.drop_box.get(0, "end"))

        if len(displayed) == 1 and displayed[0].startswith("Drop PDFs here"):
            return

        removed_paths = []
        for idx in sorted(selected, reverse=True):
            if 0 <= idx < len(self.selected_files):
                removed_paths.append(self.selected_files[idx])
                self.selected_files.pop(idx)

        self._cleanup_removed_zip_files(removed_paths)
        if not self.selected_files:
            self._cleanup_all_zip_temp()

        self.drop_box.delete(0, "end")
        if not self.selected_files:
            self.drop_box.insert("end", "Drop PDFs here, or click 'Browse'.")
            self._reset_bank_if_autodetected()
        else:
            for p in self.selected_files:
                self.drop_box.insert("end", os.path.basename(p))

        self.set_status("Removed selected item(s).")

    def _prompt_duplicate_action(self, message_text: str) -> bool:
        dialog = tk.Toplevel(self)
        dialog.title("Duplicate statements detected")
        dialog.transient(self)
        dialog.resizable(True, True)
        dialog.minsize(700, 380)

        result = {"remove": False}

        def _abort():
            result["remove"] = False
            dialog.destroy()

        def _remove_and_continue():
            result["remove"] = True
            dialog.destroy()

        dialog.protocol("WM_DELETE_WINDOW", _abort)

        body = ttk.Frame(dialog, padding=12)
        body.pack(fill="both", expand=True)

        text_frame = ttk.Frame(body)
        text_frame.pack(fill="both", expand=True)

        msg_text = tk.Text(text_frame, wrap="word", height=16)
        msg_text.pack(side="left", fill="both", expand=True)
        y_scroll = ttk.Scrollbar(text_frame, orient="vertical", command=msg_text.yview)
        y_scroll.pack(side="right", fill="y")
        msg_text.configure(yscrollcommand=y_scroll.set)
        msg_text.insert("1.0", message_text)
        msg_text.configure(state="disabled")

        btn_row = ttk.Frame(body)
        btn_row.pack(fill="x", pady=(10, 0))

        ttk.Button(btn_row, text="Remove duplicates and continue", command=_remove_and_continue).pack(side="left")
        ttk.Button(btn_row, text="OK", command=_abort).pack(side="left", padx=(8, 0))

        dialog.grab_set()
        dialog.focus_set()
        self.wait_window(dialog)
        return bool(result["remove"])

    def on_drop(self, event):
        files = parse_dnd_event_files(event.data)
        zip_paths = [p for p in files if p.lower().endswith(".zip")]
        other_paths = [p for p in files if p not in zip_paths]

        for zip_path in zip_paths:
            self.add_zip(zip_path)
        self.add_files(other_paths)

    def _on_bank_selected(self, _event=None):
        self._bank_set_by_autodetect = False

    def _reset_bank_if_autodetected(self):
        if not self.selected_files and self._bank_set_by_autodetect:
            self.bank_var.set("Select bank...")
            self._bank_set_by_autodetect = False

    def _on_app_close(self):
        self._cleanup_all_zip_temp()
        self.destroy()

    def _get_zip_temp_base_dir(self) -> str:
        if self._zip_temp_base_dir:
            return self._zip_temp_base_dir

        app_dir = os.path.dirname(os.path.abspath(__file__))
        preferred_base = os.path.join(app_dir, "_ZIP_TEMP")
        try:
            ensure_folder(preferred_base)
            self._zip_temp_base_dir = preferred_base
            return self._zip_temp_base_dir
        except Exception:
            fallback_base = os.path.join(tempfile.gettempdir(), "PDF_Converter_ZIP_TEMP")
            ensure_folder(fallback_base)
            self._zip_temp_base_dir = fallback_base
            self.set_status("ZIP temp folder fallback: using system temp.")
            return self._zip_temp_base_dir

    def _try_remove_dir_tree(self, path: str) -> None:
        if not path:
            return
        if not os.path.exists(path):
            return

        try:
            shutil.rmtree(path)
            return
        except Exception:
            pass

        try:
            os.rmdir(path)
            return
        except Exception:
            pass

        try:
            shutil.rmtree(path, ignore_errors=True)
        except Exception:
            pass

    def _cleanup_removed_zip_files(self, removed_paths: list[str]) -> None:
        candidate_dirs = set()
        for path in (removed_paths or []):
            temp_dir = self._zip_extracted_file_to_dir.pop(path, "")
            if temp_dir:
                candidate_dirs.add(temp_dir)
            try:
                if os.path.exists(path):
                    os.remove(path)
            except Exception:
                pass

        for temp_dir in candidate_dirs:
            if temp_dir not in self._zip_extracted_file_to_dir.values():
                self._try_remove_dir_tree(temp_dir)

        if not self._zip_extracted_file_to_dir:
            self._cleanup_all_zip_temp()

    def _cleanup_all_zip_temp(self) -> None:
        if self._zip_temp_base_dir:
            self._try_remove_dir_tree(self._zip_temp_base_dir)
        self._zip_extracted_file_to_dir.clear()
        self._zip_temp_base_dir = ""

    def _extract_pdfs_from_zip(self, zip_path: str) -> list[str]:
        extracted_paths: list[str] = []

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        zip_base = os.path.splitext(os.path.basename(zip_path))[0] or "zip_import"
        base = self._get_zip_temp_base_dir()
        target_dir = os.path.join(base, f"{sanitize_filename(zip_base) or 'zip_import'}_{ts}")
        ensure_folder(target_dir)

        with zipfile.ZipFile(zip_path) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue

                member_name = str(info.filename or "")
                lower_name = member_name.lower()
                if not lower_name.endswith(".pdf"):
                    continue
                if ".." in member_name.replace("\\", "/").split("/"):
                    continue
                if os.path.isabs(member_name) or member_name.startswith("/") or member_name.startswith("\\"):
                    continue

                base_name = os.path.basename(member_name)
                safe_name = sanitize_filename(base_name) or "statement.pdf"
                if not safe_name.lower().endswith(".pdf"):
                    safe_name = f"{safe_name}.pdf"

                out_path = make_unique_path(os.path.join(target_dir, safe_name))
                with zf.open(info, "r") as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted_paths.append(out_path)
                self._zip_extracted_file_to_dir[out_path] = target_dir

        if not extracted_paths:
            messagebox.showwarning("ZIP Import", "No PDF files were found in the selected ZIP.")

        return extracted_paths

    def add_zip(self, zip_path: str) -> None:
        zip_path = (zip_path or "").strip()
        if not zip_path:
            return
        if not os.path.exists(zip_path):
            self.set_status("ZIP file not found.")
            return

        try:
            extracted_paths = self._extract_pdfs_from_zip(zip_path)
        except Exception as e:
            messagebox.showerror("ZIP Import", f"Failed to import ZIP: {e}")
            return

        if extracted_paths:
            self.add_files(extracted_paths)

    def add_files(self, files: list[str]):
        pdfs = []
        for f in files:
            f = f.strip()
            if not f:
                continue
            if not os.path.exists(f):
                continue
            if not is_pdf(f):
                continue
            pdfs.append(f)

        if not pdfs:
            self.set_status("No valid PDFs added.")
            return

        if self.auto_detect_var.get():
            detected = auto_detect_bank_from_pdf(pdfs[0])
            if detected and detected in BANK_OPTIONS:
                if detected != self.bank_var.get():
                    self.bank_var.set(detected)
                    self._bank_set_by_autodetect = True

        for p in pdfs:
            if p not in self.selected_files:
                self.selected_files.append(p)

        self.drop_box.delete(0, "end")
        for p in self.selected_files:
            self.drop_box.insert("end", os.path.basename(p))

        self.set_status(f"Added {len(pdfs)} PDF(s). Total: {len(self.selected_files)}.")

    def clean_up(self):
        if not self.selected_files:
            messagebox.showwarning("No files", "Please add at least one PDF statement.")
            return

        bank = self.bank_var.get().strip()
        out_folder = self.output_folder_var.get().strip()

        if not bank or bank == "Select bank...":
            messagebox.showwarning(
                "Bank",
                "Please select a bank (or enable auto-detect and add a PDF).",
            )
            return

        try:
            self.set_status(f"Loading parser for {bank}...")
            self.set_status("Starting...")
            self.set_progress(0, max(1, len(self.selected_files)))
            parser = load_parser_module(bank)

            client_name = ""
            try:
                if hasattr(parser, "extract_account_holder_name"):
                    client_name = parser.extract_account_holder_name(self.selected_files[0]) or ""
            except Exception:
                client_name = ""
            if not client_name:
                client_name = get_client_name_from_pdf(self.selected_files[0])

            client_folder = sanitize_filename((client_name or "").strip().upper()) or "CLIENT"

            initial_dir = ""
            if out_folder:
                initial_dir = out_folder
            else:
                try:
                    initial_dir = os.path.dirname(self.selected_files[0])
                except Exception:
                    initial_dir = ""

            zip_path = filedialog.asksaveasfilename(
                title="Save ZIP file",
                defaultextension=".zip",
                filetypes=[("ZIP file", "*.zip")],
                initialdir=initial_dir or None,
                initialfile=f"{client_folder}.zip",
            )

            if not zip_path:
                self.set_status("Clean Up cancelled.")
                return

            try:
                ensure_folder(os.path.dirname(zip_path))
            except Exception as e:
                messagebox.showerror("Save error", f"Cannot create folder for ZIP file:\n{e}")
                self.set_status("Error.")
                return

            try:
                self.output_folder_var.set(os.path.dirname(zip_path))
            except Exception:
                pass

            items = []
            failures = []

            def _coerce_date(value):
                try:
                    if hasattr(value, "to_pydatetime"):
                        value = value.to_pydatetime()
                except Exception:
                    return None
                if isinstance(value, datetime):
                    return value.date()
                if isinstance(value, date):
                    return value
                return None

            def _get_period_dates(pdf_path):
                dmin = None
                dmax = None

                if callable(getattr(parser, "extract_statement_period", None)):
                    try:
                        period = parser.extract_statement_period(pdf_path)
                        if isinstance(period, (tuple, list)) and len(period) >= 2:
                            pstart = _coerce_date(period[0])
                            pend = _coerce_date(period[1])
                            if pstart and pend:
                                dmin, dmax = pstart, pend
                    except Exception:
                        dmin, dmax = None, None

                if not (dmin and dmax):
                    try:
                        txns = parser.extract_transactions(pdf_path) or []
                        _dates = [t.get("Date") for t in (txns or []) if t.get("Date")]
                        dmin = min(_dates) if _dates else None
                        dmax = max(_dates) if _dates else None
                    except Exception:
                        failures.append(os.path.basename(pdf_path))
                        dmin, dmax = None, None

                period_str = ""
                try:
                    if dmin and dmax:
                        period_str = f"{dmin.strftime('%d.%m.%y')} - {dmax.strftime('%d.%m.%y')}"
                except Exception:
                    period_str = ""

                return dmin, dmax, period_str

            for i, pdf_path in enumerate(self.selected_files, start=1):
                self.set_status(
                    f"Reading statement dates {i}/{len(self.selected_files)}: {os.path.basename(pdf_path)}"
                )
                dmin, dmax, period = _get_period_dates(pdf_path)

                items.append(
                    {
                        "path": pdf_path,
                        "pdf": os.path.basename(pdf_path),
                        "date_min": dmin,
                        "date_max": dmax,
                        "period": period,
                    }
                )

            def _sort_key(it):
                d = it.get("date_min")
                try:
                    if hasattr(d, "to_pydatetime"):
                        d = d.to_pydatetime()
                except Exception:
                    pass
                try:
                    if isinstance(d, datetime):
                        d = d.date()
                except Exception:
                    pass

                if d is None:
                    return (1, datetime.max.date(), it.get("pdf", ""))
                return (0, d, it.get("pdf", ""))

            items.sort(key=_sort_key)

            with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                for idx, it in enumerate(items, start=1):
                    period = it.get("period") or ""
                    if period:
                        arc_base = f"{idx} {period}.pdf"
                    else:
                        base = sanitize_filename(os.path.splitext(it.get("pdf", "statement"))[0]) or "statement"
                        arc_base = f"{idx} {base}.pdf"

                    arcname = arc_base
                    zf.write(it["path"], arcname=arcname)

            self.set_status(f"Clean Up complete: {zip_path}")

            if failures:
                messagebox.showwarning(
                    "Clean Up complete (some periods unknown)",
                    "ZIP created successfully, but I could not extract dates for some PDFs.\n"
                    "They were included using the original filename as a fallback:\n\n"
                    + "\n".join(failures)
                    + f"\n\nZIP: {zip_path}",
                )
            else:
                messagebox.showinfo("Clean Up complete", f"ZIP created:\n{zip_path}")

        except Exception as e:
            self.set_status("Error.")
            err = "".join(traceback.format_exception(type(e), e, e.__traceback__))
            messagebox.showerror("Clean Up error", f"{e}\n\nDetails:\n{err}")

    # (run_parser unchanged)
    def run_parser(self):
        if not self.selected_files:
            messagebox.showwarning("No files", "Please add at least one PDF statement.")
            return

        bank = self.bank_var.get().strip()
        out_folder = self.output_folder_var.get().strip()

        if not bank or bank == "Select bank...":
            messagebox.showwarning(
                "Bank",
                "Please select a bank (or enable auto-detect and add a PDF).",
            )
            return
        if out_folder:
            try:
                ensure_folder(out_folder)
            except Exception:
                out_folder = ""

        try:
            self.set_status(f"Loading parser for {bank}...")
            parser = load_parser_module(bank)

            client_name = ""
            try:
                if hasattr(parser, "extract_account_holder_name"):
                    client_name = parser.extract_account_holder_name(self.selected_files[0]) or ""
            except Exception:
                client_name = ""
            if not client_name:
                client_name = get_client_name_from_pdf(self.selected_files[0])

            all_transactions = []
            recon_results = []
            audit_results = []
            per_pdf_txns: dict[str, list[dict]] = {}
            remove_txn_ids: set[int] = set()
            pdf_by_name: dict[str, str] = {os.path.basename(p): p for p in (self.selected_files or [])}

            if self.auto_detect_var.get() and len(self.selected_files) > 1:
                self.set_status("Detecting bank...")
                mismatches = []
                unknowns = []

                for p in self.selected_files:
                    detected = auto_detect_bank_from_pdf(p)
                    if detected is None:
                        unknowns.append(os.path.basename(p))
                    elif detected != bank:
                        mismatches.append(f"{os.path.basename(p)} → {detected}")

                if mismatches:
                    NL = chr(10)
                    msg = (
                        "Auto-detect thinks some PDFs may be from a different bank than the one selected." + NL + NL
                        + f"Selected bank: {bank}" + NL + NL
                        + NL.join(mismatches)
                        + NL + NL
                        + "This can happen if a statement contains another bank name/BIC in a transaction description, "
                        "or if the PDF has a cover/summary page." + NL + NL
                        + f"Do you want to continue and parse ALL PDFs as {bank}?"
                    )

                    proceed_anyway = messagebox.askyesno("Bank not confirmed", msg)
                    if not proceed_anyway:
                        self.set_status("Cancelled: bank not confirmed.")
                        return

                if unknowns:
                    NL = chr(10)
                    msg = (
                        "Auto-detect could not confirm the bank for some PDFs." + NL
                        + "I will continue, but results may be wrong if a different bank is included:" + NL + NL
                        + NL.join(unknowns)
                    )
                    messagebox.showwarning("Bank not confirmed", msg)

            def _status_startswith(v, prefix: str) -> bool:
                try:
                    return str(v or "").startswith(prefix)
                except Exception:
                    return False

            def _get_pdf_period(_parser, _pdf_path: str, _rec: dict):
                """Best-effort (period_start, period_end) from parser; else None/None.

                Expected parser hook (if present): extract_statement_period(pdf_path) -> (date|None, date|None)
                Also tries common alternative names.
                """
                fn_names = [
                    "extract_statement_period",
                    "extract_statement_period_dates",
                    "extract_period",
                    "extract_statement_date_range",
                    "extract_date_range",
                ]
                for nm in fn_names:
                    try:
                        fn = getattr(_parser, nm, None)
                    except Exception:
                        fn = None
                    if not callable(fn):
                        continue
                    try:
                        ps, pe = fn(_pdf_path)
                        return ps, pe
                    except TypeError:
                        # Some implementations may take (path, rec)
                        try:
                            ps, pe = fn(_pdf_path, _rec)
                            return ps, pe
                        except Exception:
                            continue
                    except Exception:
                        continue
                return None, None

            run_log_lines = []
            run_log_lines.append(f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            run_log_lines.append(f"Bank: {bank}")
            run_log_lines.append(f"PDF count: {len(self.selected_files)}")
            run_log_lines.append("")

            for i, pdf_path in enumerate(self.selected_files, start=1):
                self.set_status(f"Parsing: {os.path.basename(pdf_path)} ({i}/{len(self.selected_files)})")
                self.set_progress(i, max(1, len(self.selected_files)))
                txns = parser.extract_transactions(pdf_path)
                per_pdf_txns[pdf_path] = txns
                all_transactions.extend(txns)
                rec = reconcile_statement(parser, pdf_path, txns)
                audit = run_audit_checks_basic(
                    os.path.basename(pdf_path),
                    txns,
                    rec.get("start_balance"),
                    rec.get("end_balance"),
                )

                # Provide txns to core continuity logic (overlap resolution uses these when present).
                rec["transactions"] = txns
                try:
                    _dates = [t.get("Date") for t in (txns or []) if t.get("Date")]
                    rec["date_min"] = min(_dates) if _dates else None
                    rec["date_max"] = max(_dates) if _dates else None
                except Exception:
                    rec["date_min"] = None
                    rec["date_max"] = None

                # Capture per-PDF statement period ONLY if parser exposes it; otherwise leave None.
                # (Starling and others may provide: extract_statement_period(pdf_path) -> (start, end))
                try:
                    if hasattr(parser, "extract_statement_period") and callable(getattr(parser, "extract_statement_period")):
                        ps, pe = parser.extract_statement_period(pdf_path)
                    else:
                        ps, pe = (None, None)
                except Exception:
                    ps, pe = (None, None)

                if hasattr(ps, "to_pydatetime"):
                    try:
                        ps = ps.to_pydatetime()
                    except Exception:
                        pass
                if hasattr(pe, "to_pydatetime"):
                    try:
                        pe = pe.to_pydatetime()
                    except Exception:
                        pass

                if isinstance(ps, (date, datetime)):
                    rec["period_start"] = ps
                elif rec.get("period_start") is None:
                    rec["period_start"] = None

                if isinstance(pe, (date, datetime)):
                    rec["period_end"] = pe
                elif rec.get("period_end") is None:
                    rec["period_end"] = None

                if bank == "Lloyds":
                    try:
                        start_d = rec.get("date_min")
                        opening = None
                        if start_d:
                            first_balance_idx = None
                            for j, t in enumerate(txns or []):
                                if t.get("Date") != start_d:
                                    continue
                                bal = t.get("Balance")
                                if bal is None or bal == "":
                                    continue
                                first_balance_idx = j
                                break

                            if first_balance_idx is not None:
                                bal_val = float((txns[first_balance_idx] or {}).get("Balance"))
                                net = 0.0
                                for j in range(0, first_balance_idx + 1):
                                    tt = txns[j]
                                    if tt.get("Date") == start_d:
                                        amt = tt.get("Amount")
                                        if amt is None or amt == "":
                                            continue
                                        net += float(amt)
                                opening = round(bal_val - net, 2)

                        rec["continuity_start_balance"] = opening if opening is not None else rec.get("start_balance")
                    except Exception:
                        rec["continuity_start_balance"] = rec.get("start_balance")
                else:
                    rec["continuity_start_balance"] = rec.get("start_balance")

                try:
                    rec["txn_count"] = len(txns or [])
                except Exception:
                    rec["txn_count"] = None
                try:
                    rec["fingerprint"] = compute_statement_fingerprint(txns)
                except Exception:
                    rec["fingerprint"] = None
                rec["pdf_path"] = pdf_path

                recon_results.append(rec)
                audit_results.append(audit)

                run_log_lines.append(f"{os.path.basename(pdf_path)}")
                run_log_lines.append(f"  Transactions: {len(txns)}")
                run_log_lines.append(f"  Reconciliation: {rec.get('status')}")

                # Period visibility for debugging (core overlap/chronology uses period_start/period_end).
                try:
                    ps = rec.get("period_start")
                    pe = rec.get("period_end")
                    if ps and pe and hasattr(ps, "strftime") and hasattr(pe, "strftime"):
                        run_log_lines.append(f"  Period: {ps.strftime('%d/%m/%Y')} - {pe.strftime('%d/%m/%Y')}")
                    else:
                        run_log_lines.append("  Period: None")
                except Exception:
                    run_log_lines.append("  Period: None")

                if rec.get("status") in ("OK", "Mismatch"):
                    run_log_lines.append(
                        f"  Start: {_fmt_money(rec.get('start_balance'))} | Net: {_fmt_money(rec.get('sum_amounts'))} | End: {_fmt_money(rec.get('end_balance'))}"
                    )
                    if rec.get("status") == "Mismatch":
                        run_log_lines.append(f"  Diff: {_fmt_money(rec.get('difference'))}")
                run_log_lines.append(
                    f"  Balance Walk: {audit.get('balance_walk_status')} | {audit.get('balance_walk_summary') or ''}"
                )
                run_log_lines.append(
                    f"  Row Shape Sanity: {audit.get('row_shape_status')} | {audit.get('row_shape_summary') or ''}"
                )
                run_log_lines.append("")

            self.set_status("Running reconciliation checks...")
            duplicate_groups = find_duplicate_statements(recon_results)
            if duplicate_groups:
                def _fmt_date(v):
                    try:
                        if v is None or v == "":
                            return ""
                        if hasattr(v, "to_pydatetime"):
                            v = v.to_pydatetime()
                        if hasattr(v, "strftime"):
                            return v.strftime("%d/%m/%Y")
                        return str(v)
                    except Exception:
                        return str(v)

                msg_lines = []
                msg_lines.append("Duplicate statements detected. Please remove the duplicates and run again.")
                msg_lines.append("")

                for gi, grp in enumerate(duplicate_groups, start=1):
                    msg_lines.append(f"Group {gi} ({len(grp)} files):")
                    for r in grp:
                        msg_lines.append(f"  - {r.get('pdf')}")

                    try:
                        dmin = _fmt_date(grp[0].get("date_min"))
                        dmax = _fmt_date(grp[0].get("date_max"))
                        dr = (f"{dmin} - {dmax}").strip(" -")
                    except Exception:
                        dr = ""

                    try:
                        start = _fmt_money(grp[0].get("start_balance"))
                        end = _fmt_money(grp[0].get("end_balance"))
                    except Exception:
                        start, end = "", ""

                    try:
                        txc = grp[0].get("txn_count")
                    except Exception:
                        txc = ""

                    summary_parts = []
                    if dr:
                        summary_parts.append(f"Dates {dr}")
                    if start or end:
                        summary_parts.append(f"Start {start} / End {end}")
                    if txc not in (None, ""):
                        summary_parts.append(f"Txns {txc}")

                    if summary_parts:
                        msg_lines.append("  Summary: " + " | ".join(summary_parts))

                    msg_lines.append("")

                try:
                    ensure_folder(LOGS_DIR)
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    log_path = os.path.join(LOGS_DIR, f"duplicate_statements_{ts}.txt")
                    with open(log_path, "w", encoding="utf-8") as f:
                        f.write("\n".join(msg_lines).strip() + "\n")
                except Exception:
                    pass

                should_remove_duplicates = self._prompt_duplicate_action("\n".join(msg_lines).strip())
                if not should_remove_duplicates:
                    self.set_status("Error: duplicate statements detected.")
                    return

                original_selected_files = list(self.selected_files or [])

                def _resolve_pdf_path(_rec, _selected_files, _used_paths=None):
                    _used_paths = _used_paths or set()
                    path = _rec.get("pdf_path")
                    if path:
                        return path
                    name = os.path.basename(str(_rec.get("pdf") or ""))
                    if not name:
                        return ""
                    matches = [p for p in (_selected_files or []) if os.path.basename(p) == name]
                    if not matches:
                        return ""
                    for m in matches:
                        if m not in _used_paths:
                            return m
                    return matches[0]

                removed_pdf_paths: set[str] = set()
                kept_pdf_paths: set[str] = set()
                for grp in duplicate_groups:
                    if not grp:
                        continue
                    assigned_group_paths: set[str] = set()
                    keep_path = _resolve_pdf_path(grp[0], original_selected_files, assigned_group_paths)
                    if keep_path:
                        kept_pdf_paths.add(keep_path)
                        assigned_group_paths.add(keep_path)

                    for dup_rec in grp[1:]:
                        dup_path = _resolve_pdf_path(dup_rec, original_selected_files, assigned_group_paths)
                        if dup_path:
                            removed_pdf_paths.add(dup_path)
                            assigned_group_paths.add(dup_path)

                removed_pdf_paths = removed_pdf_paths - kept_pdf_paths

                self.selected_files = [p for p in self.selected_files if p not in removed_pdf_paths]
                self._cleanup_removed_zip_files(list(removed_pdf_paths))

                for removed_path in removed_pdf_paths:
                    per_pdf_txns.pop(removed_path, None)

                filtered_pairs = []
                for rec, audit in zip(recon_results, audit_results):
                    rec_path = _resolve_pdf_path(rec, original_selected_files)
                    if rec_path and rec_path in removed_pdf_paths:
                        continue
                    filtered_pairs.append((rec, audit))
                recon_results = [pair[0] for pair in filtered_pairs]
                audit_results = [pair[1] for pair in filtered_pairs]

                all_transactions = []
                for p in self.selected_files:
                    all_transactions.extend(per_pdf_txns.get(p) or [])

                self.drop_box.delete(0, "end")
                if not self.selected_files:
                    self.drop_box.insert("end", "Drop PDFs here, or click 'Browse'.")
                else:
                    for p in self.selected_files:
                        self.drop_box.insert("end", os.path.basename(p))

                self.set_status("Duplicates removed. Continuing...")

            if not all_transactions:
                log_text = "\n".join(run_log_lines).rstrip() + "\n"

                initial_dir = ""
                if out_folder:
                    initial_dir = out_folder
                else:
                    try:
                        initial_dir = os.path.dirname(self.selected_files[0])
                    except Exception:
                        initial_dir = ""

                autodetect_first_pdf = None
                try:
                    if self.auto_detect_var.get() and self.selected_files:
                        autodetect_first_pdf = auto_detect_bank_from_pdf(self.selected_files[0])
                except Exception:
                    autodetect_first_pdf = None

                parser_file = ""
                try:
                    parser_file = getattr(parser, "__file__", "") or ""
                except Exception:
                    parser_file = ""

                bundle_base = sanitize_filename(f"{client_name or 'RUN'} - No Transactions") or "RUN - No Transactions"

                self.last_saved_output_path = None
                self.last_report_data = {
                    "recon_results": recon_results,
                    "audit_results": audit_results,
                    "continuity_results": [],
                    "coverage_period": "",
                    "source_pdfs": list(self.selected_files or []),
                    "any_warn": True,
                    "log_path": "",
                    "log_text": log_text,
                    "learning_report_path": None,
                    "learning_report_inline": "",
                    "learning_report_error": "",
                    "learning_report_generated": False,
                    "output_xlsx_path": None,
                    "bundle_base": bundle_base,
                    "bank": bank,
                    "autodetect_first_pdf": autodetect_first_pdf,
                    "parser_file": parser_file,
                    "client_name": client_name,
                    "run_filename": "",
                }
                self.last_excel_data = {
                    "transactions": [],
                    "client_name": client_name,
                    "filename": "No Transactions Found.xlsx",
                    "initial_dir": initial_dir,
                }

                try:
                    report_path, report_text, report_err = self.generate_learning_report(
                        reason="No transactions found", write_to_disk=False
                    )
                    if self.last_report_data is not None:
                        self.last_report_data["learning_report_path"] = report_path
                        self.last_report_data["learning_report_inline"] = report_text or ""
                        self.last_report_data["learning_report_error"] = report_err or ""
                except Exception:
                    pass

                self.create_support_bundle_zip()

                show_reconciliation_popup(
                    self,
                    "(Not saved yet)",
                    recon_results,
                    coverage_period="",
                    continuity_results=[],
                    audit_results=audit_results,
                    pre_save=True,
                    open_log_folder_callback=self.open_log_folder,
                    client_name=client_name,
                )

                messagebox.showwarning(
                    "No transactions found",
                    "No transactions were extracted from the selected PDFs.\n\n"
                    "Likely causes:\n"
                    "• The wrong bank parser is selected\n"
                    "• The statement format is new/unsupported\n"
                    "• The PDF is scanned/image-only and has no extractable text\n\n"
                    "Use Open Log and send the automatically created support bundle ZIP for investigation.",
                )
                self.set_status("Done with warnings. No transactions found.")
                return

            def _date_key(x):
                d = x.get("Date")
                return d if d is not None else datetime.min.date()

            all_transactions.sort(key=_date_key)

            date_values = []
            for t in all_transactions:
                d = t.get("Date")
                if d:
                    date_values.append(d)
            date_min = min(date_values) if date_values else None
            date_max = max(date_values) if date_values else None

            def _coerce_to_date(v):
                try:
                    if v is None or v == "":
                        return None
                    if hasattr(v, "to_pydatetime"):
                        v = v.to_pydatetime()
                    if isinstance(v, datetime):
                        return v.date()
                    if isinstance(v, date):
                        return v
                    if hasattr(v, "date"):
                        dv = v.date()
                        if isinstance(dv, date):
                            return dv
                except Exception:
                    return None
                return None

            period_starts = []
            period_ends = []
            for rec in (recon_results or []):
                ps = _coerce_to_date(rec.get("period_start"))
                pe = _coerce_to_date(rec.get("period_end"))
                if ps and pe:
                    period_starts.append(ps)
                    period_ends.append(pe)

            statement_period_start = min(period_starts) if period_starts else None
            statement_period_end = max(period_ends) if period_ends else None

            if statement_period_start and statement_period_end:
                filename = build_output_filename(client_name, statement_period_start, statement_period_end)
            else:
                filename = build_output_filename(client_name, date_min, date_max)

            self.set_status("Running continuity checks...")
            continuity_results = compute_statement_continuity(recon_results)

            # Apply overlap de-duplication results produced by core continuity logic.
            # Core will populate overlap_* fields on each continuity link when applicable.
            rec_by_pdfname: dict[str, dict] = {}
            try:
                for r in (recon_results or []):
                    nm = r.get("pdf")
                    if nm:
                        rec_by_pdfname[str(nm)] = r
            except Exception:
                rec_by_pdfname = {}

            for link in (continuity_results or []):
                # Prefer core's display_status for UI/popup/logging.
                try:
                    ds = link.get("display_status")
                    if ds:
                        link["status"] = str(ds)
                except Exception:
                    pass

                prev_pdf = str(link.get("prev_pdf") or "")
                next_pdf = str(link.get("next_pdf") or "")

                applied = False
                try:
                    applied = bool(link.get("applied_overlap_resolution"))
                except Exception:
                    applied = False

                dup_idx = link.get("duplicates_to_remove_from_B")
                if not isinstance(dup_idx, list):
                    dup_idx = []

                removed_n_effective = 0
                expected_removed_for_link = 0

                if applied and dup_idx:
                    # Apply removal plan: remove indices from B ONLY, in descending order.
                    # IMPORTANT: ensure we remove from the SAME B txn list that will be used to rebuild the combined export.
                    b_rec = rec_by_pdfname.get(next_pdf)
                    b_txns = None
                    try:
                        if b_rec is not None:
                            b_txns = b_rec.get("transactions")
                    except Exception:
                        b_txns = None

                    if not isinstance(b_txns, list):
                        b_txns = []

                    # Locate B's original pdf_path so we can also mutate per_pdf_txns (if it is a different list object).
                    b_path = None
                    try:
                        b_path = pdf_by_name.get(next_pdf)
                    except Exception:
                        b_path = None
                    if not b_path:
                        try:
                            for p in (self.selected_files or []):
                                if os.path.basename(p) == next_pdf:
                                    b_path = p
                                    break
                        except Exception:
                            b_path = None

                    b_path_txns = None
                    try:
                        if b_path:
                            b_path_txns = per_pdf_txns.get(b_path)
                    except Exception:
                        b_path_txns = None

                    if not isinstance(b_path_txns, list):
                        b_path_txns = None

                    # Normalise & sort indices descending to avoid shifting.
                    idxs = []
                    for x in dup_idx:
                        try:
                            idxs.append(int(x))
                        except Exception:
                            continue
                    idxs = sorted(set(idxs), reverse=True)

                    expected_removed_for_link = len(idxs)

                    # Remove from b_txns
                    for ii in idxs:
                        if 0 <= ii < len(b_txns):
                            try:
                                removed_txn = b_txns.pop(ii)
                                remove_txn_ids.add(id(removed_txn))  # fallback safety
                                removed_n_effective += 1
                            except Exception:
                                pass

                    # If per_pdf_txns uses a different list object, remove there too so rebuild matches.
                    if b_path_txns is not None and b_path_txns is not b_txns:
                        for ii in idxs:
                            if 0 <= ii < len(b_path_txns):
                                try:
                                    removed_txn = b_path_txns.pop(ii)
                                    remove_txn_ids.add(id(removed_txn))
                                except Exception:
                                    pass

                # Logging / verification for each continuity link
                # Track expected vs effective removals for output verification.
                try:
                    expected_total_removed_from_core = expected_total_removed_from_core
                except Exception:
                    expected_total_removed_from_core = 0
                try:
                    effective_total_removed_in_lists = effective_total_removed_in_lists
                except Exception:
                    effective_total_removed_in_lists = 0

                if applied and dup_idx:
                    try:
                        expected_total_removed_from_core += int(link.get("removed_count") or expected_removed_for_link or 0)
                    except Exception:
                        expected_total_removed_from_core += int(expected_removed_for_link or 0)
                    effective_total_removed_in_lists += int(removed_n_effective or 0)

                try:
                    ow = link.get("overlap_window")
                    win = ""
                    if isinstance(ow, dict):
                        ws = ow.get("start")
                        we = ow.get("end")
                        if ws and we and hasattr(ws, "strftime") and hasattr(we, "strftime"):
                            win = f"{ws.strftime('%d/%m/%Y')} - {we.strftime('%d/%m/%Y')}"
                except Exception:
                    win = ""

                removed_count = link.get("removed_count")
                if removed_count in (None, ""):
                    removed_count = removed_n_effective
                dupe_sum = link.get("dupe_sum")

                chrono_applied = link.get("chronology_gate_applied")
                chrono_note = link.get("chronology_gate_note")

                if applied:
                    run_log_lines.append(
                        f"Continuity: {prev_pdf} -> {next_pdf} | overlap YES"
                        + (f" | window {win}" if win else "")
                        + f" | removed {removed_count}"
                        + (f" | dupe_sum {_fmt_money(dupe_sum)}" if dupe_sum not in (None, "") else "")
                    )
                else:
                    first_overlap_line = ""
                    try:
                        lines = link.get("overlap_log_lines")
                        if isinstance(lines, list) and lines:
                            first_overlap_line = str(lines[0])
                    except Exception:
                        first_overlap_line = ""

                    run_log_lines.append(
                        f"Continuity: {prev_pdf} -> {next_pdf} | overlap NO"
                        + (f" | note {first_overlap_line}" if first_overlap_line else "")
                    )

                # Chain/chronology debug info (lightweight)
                try:
                    if chrono_applied is not None:
                        run_log_lines.append(
                            f"  Chronology gate applied: {'YES' if chrono_applied else 'NO'}"
                            + (f" | {chrono_note}" if chrono_note else "")
                        )
                except Exception:
                    pass

                try:
                    cc_total = link.get("chain_candidates_total")
                    if cc_total is not None:
                        run_log_lines.append(
                            "  Chain candidates: "
                            + f"total={link.get('chain_candidates_total')} "
                            + f"known={link.get('chain_candidates_known_period_start')} "
                            + f"pass={link.get('chain_candidates_chrono_pass')} "
                            + f"fail={link.get('chain_candidates_chrono_fail')} "
                            + f"unknown={link.get('chain_candidates_chrono_unknown')}"
                        )
                except Exception:
                    pass

                run_log_lines.append("")

            # Rebuild combined transactions AFTER applying de-dupe (so Excel output reflects the plan).
            # Rebuild whenever any overlap removal was applied, even if ids couldn't be captured.
            try:
                _need_rebuild = bool(effective_total_removed_in_lists)
            except Exception:
                _need_rebuild = bool(remove_txn_ids)

            if _need_rebuild:
                rebuilt = []
                for p in (self.selected_files or []):
                    rebuilt.extend(per_pdf_txns.get(p) or [])

                all_transactions = list(rebuilt)

                # Log verification: expected (core) vs effective removals vs output delta.
                try:
                    expected_n = int(expected_total_removed_from_core)
                except Exception:
                    expected_n = 0
                try:
                    effective_n = int(effective_total_removed_in_lists)
                except Exception:
                    effective_n = 0

                run_log_lines.append(
                    f"Overlap de-duplication applied: removed {effective_n} transactions from output"
                    + (f" (core expected {expected_n})" if expected_n else "")
                )
                run_log_lines.append("")

            any_gap = any(_status_startswith((r.get('status') or ''), 'Mismatch') for r in (continuity_results or []))

            coverage_period = ""
            try:
                if date_min and date_max:
                    coverage_period = f"{date_min.strftime('%d/%m/%Y')} to {date_max.strftime('%d/%m/%Y')}"
            except Exception:
                coverage_period = ""
            any_issue = any(
                (r.get("status") not in ("OK", "Not checked"))
                for r in (recon_results or [])
            )

            # Continuity: anything not starting with OK (including NOT CHECKED / balances not found)
            # should be treated as an issue so logs/support bundles are produced.
            any_cont_issue = any(
                not _status_startswith((c.get("status") or c.get("display_status") or ""), "OK")
                for c in (continuity_results or [])
            )

            any_issue = any_issue or any_cont_issue or any_gap
            any_audit_issue = any((a.get("status") or "") != "OK" for a in (audit_results or []))
            any_issue = any_issue or any_audit_issue

            log_text = "\n".join(run_log_lines).rstrip() + "\n"
            if self.last_report_data is None:
                self.last_report_data = {}
            self.last_report_data["log_text"] = log_text

            all_ok = all((r.get("status") == "OK") for r in recon_results)

            def _is_ok_status(s: str) -> bool:
                try:
                    return str(s or "").strip().upper().startswith("OK")
                except Exception:
                    return False

            if not continuity_results:
                cont_ok = False
            else:
                cont_ok = True
                for link in continuity_results:
                    if not isinstance(link, dict):
                        cont_ok = False
                        break
                    st = link.get("display_status") or link.get("status") or ""
                    if not _is_ok_status(st):
                        cont_ok = False
                        break

            full_pass = all_ok and cont_ok
            audit_ok = all((a.get("status") == "OK") for a in (audit_results or []))
            full_pass = all_ok and cont_ok and audit_ok

            if not full_pass:
                try:
                    ensure_folder(LOGS_DIR)
                    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                    base = sanitize_filename(os.path.splitext(filename)[0]) or "RUN"
                    recon_log_path = make_unique_path(
                        os.path.join(LOGS_DIR, f"{base} - recon log - {ts}.txt")
                    )
                    with open(recon_log_path, "w", encoding="utf-8") as f:
                        f.write(log_text)
                except Exception:
                    recon_log_path = None
            else:
                recon_log_path = ""

            initial_dir = ""
            if out_folder:
                initial_dir = out_folder
            else:
                try:
                    initial_dir = os.path.dirname(self.selected_files[0])
                except Exception:
                    initial_dir = ""

            self.last_saved_output_path = None
            any_warn = any((r.get("status") or "") != "OK" for r in (recon_results or [])) or any(
                not _status_startswith((c.get("display_status") or c.get("status") or ""), "OK")
                for c in (continuity_results or [])
                if isinstance(c, dict)
            )
            any_audit_warn = any((a.get("status") or "") != "OK" for a in (audit_results or []))
            any_warn = any_warn or any_audit_warn

            autodetect_first_pdf = None
            try:
                if self.auto_detect_var.get() and self.selected_files:
                    autodetect_first_pdf = auto_detect_bank_from_pdf(self.selected_files[0])
            except Exception:
                autodetect_first_pdf = None

            parser_file = ""
            try:
                parser_file = getattr(parser, "__file__", "") or ""
            except Exception:
                parser_file = ""

            self.last_report_data = {
                "recon_results": recon_results,
                "audit_results": audit_results,
                "continuity_results": continuity_results,
                "coverage_period": coverage_period,
                "source_pdfs": list(self.selected_files or []),
                "any_warn": bool(any_warn),
                "log_path": recon_log_path,
                "log_text": log_text,
                "learning_report_path": None,
                "learning_report_inline": "",
                "learning_report_error": "",
                "learning_report_generated": False,
                "output_xlsx_path": None,
                "bundle_base": os.path.splitext(filename)[0],
                "bank": bank,
                "autodetect_first_pdf": autodetect_first_pdf,
                "parser_file": parser_file,
                "client_name": client_name,
                "run_filename": filename,
            }
            self.last_excel_data = {
                "transactions": all_transactions,
                "client_name": client_name,
                "filename": filename,
                "initial_dir": initial_dir,
                "statement_period_start": statement_period_start,
                "statement_period_end": statement_period_end,
            }

            # Auto-create a support bundle zip whenever reconciliation or continuity has warnings/errors.
            if any_warn or any_issue:
                self.create_support_bundle_zip()

            if any_issue:
                any_recon_mismatch = any((r.get("status") or "") == "Mismatch" for r in (recon_results or []))
                any_cont_mismatch = any(
                    _status_startswith((c.get("display_status") or c.get("status") or ""), "Mismatch")
                    for c in (continuity_results or [])
                    if isinstance(c, dict)
                )
                issue_reason = "Mismatch" if (any_recon_mismatch or any_cont_mismatch) else "Issue"
                try:
                    report_path, report_text, report_err = self.generate_learning_report(
                        reason=issue_reason, write_to_disk=False
                    )
                    try:
                        if self.last_report_data is not None:
                            self.last_report_data["learning_report_inline"] = report_text or ""
                            self.last_report_data["learning_report_error"] = report_err or ""
                    except Exception:
                        pass
                except Exception:
                    pass

            show_reconciliation_popup(
                self,
                "(Not saved yet)",
                recon_results,
                coverage_period=coverage_period,
                continuity_results=continuity_results,
                audit_results=audit_results,
                pre_save=True,
                open_log_folder_callback=self.open_log_folder,
                client_name=client_name,
            )

            output_path = filedialog.asksaveasfilename(
                title="Save Excel file",
                defaultextension=".xlsx",
                filetypes=[("Excel Workbook", "*.xlsx")],
                initialdir=initial_dir or None,
                initialfile=filename,
            )

            if not output_path:
                self.set_status("Cancelled.")
                return

            try:
                ensure_folder(os.path.dirname(output_path))
            except Exception as e:
                messagebox.showerror("Save error", f"Cannot create folder for output file:\n{e}")
                self.set_status("Error.")
                return

            try:
                self.output_folder_var.set(os.path.dirname(output_path))
            except Exception:
                pass

            self.set_status("Writing Excel...")
            save_transactions_to_excel(
                all_transactions,
                output_path,
                client_name=client_name,
                header_period_start=statement_period_start,
                header_period_end=statement_period_end,
            )

            self.last_saved_output_path = output_path

            try:
                if self.last_report_data is not None:
                    self.last_report_data["output_xlsx_path"] = output_path
            except Exception:
                pass

            self.set_progress(len(self.selected_files), max(1, len(self.selected_files)))
            if any_warn:
                self.set_status(f"Done with warnings. Output: {output_path}")
            else:
                self.set_status(f"Done. Output: {output_path}")

        except Exception as e:
            self.set_status("Error.")
            err = "".join(traceback.format_exception(type(e), e, e.__traceback__))

            try:
                if self.last_report_data is None:
                    self.last_report_data = {
                        "recon_results": [],
                        "audit_results": [],
                        "continuity_results": [],
                        "coverage_period": "",
                        "source_pdfs": list(self.selected_files or []),
                        "any_warn": True,
                        "log_path": None,
                        "learning_report_path": None,
                        "learning_report_inline": "",
                        "learning_report_error": "",
                        "learning_report_generated": False,
                        "output_xlsx_path": None,
                        "bundle_base": "RUN",
                        "bank": bank,
                        "autodetect_first_pdf": None,
                        "parser_file": "",
                        "client_name": "",
                        "run_filename": "",
                    }
                report_path, report_text, report_err = self.generate_learning_report(
                    reason="Exception", exception=e, write_to_disk=False
                )
                try:
                    if self.last_report_data is not None:
                        self.last_report_data["learning_report_inline"] = report_text or ""
                        self.last_report_data["learning_report_error"] = report_err or ""
                except Exception:
                    pass
            except Exception:
                pass

            try:
                ensure_folder(LOGS_DIR)
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                crash_name = f"crash_{ts}.txt"
                crash_path = os.path.join(LOGS_DIR, crash_name)
                with open(crash_path, "w", encoding="utf-8") as f:
                    f.write(f"Time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                    f.write(f"Bank: {bank}\n")
                    f.write(f"Output folder: {out_folder}\n")
                    f.write("PDFs:\n")
                    for p in (self.selected_files or []):
                        f.write(f"  - {p}\n")
                    f.write("\nException:\n")
                    f.write(err)
            except Exception:
                pass

            messagebox.showerror("Error", f"{e}\n\nDetails:\n{err}")


def _self_tests():
    # Minimal, optional sanity checks (no GUI launched).
    assert _fmt_money(None) == ""
    assert _fmt_money(0) == "£0.00"
    assert _fmt_money(12.3) == "£12.30"
    assert _fmt_money(-12.3) == "-£12.30"
    assert _fmt_money("£1,234.50") == "£1,234.50"


if __name__ == "__main__":
    # Set GUI_SELFTEST=1 to run quick format tests without launching the app.
    if os.environ.get("GUI_SELFTEST") == "1":
        _self_tests()
    else:
        app = App()
        app.mainloop()
